#!/usr/bin/env python3
"""Cloud Security Assessment runner for the CloudFormation template baseline.

Scans every CloudFormation template (and any Terraform sources, if present) with
Checkov, cfn_nag (when installed) and cfn-lint, adds a custom rule pack for the
Government security baseline, normalizes every result into one finding schema,
maps each finding to NIST SP 800-53 Rev. 5 controls, CIS AWS Foundations
Benchmark v3.0.0 recommendations, a DoD Cloud Computing SRG topic area and a
DISA CAT severity, and writes:

  security-assessment/findings.json
  security-assessment/Cloud-Security-Findings-Tracker.xlsx

Usage:
  python3 scripts/cloud_security_assessment.py [--repo-root .] [--out-dir security-assessment]
        [--skip-checkov] [--skip-cfn-nag] [--skip-cfn-lint]
        [--baseline security-assessment/findings.json]   # previous run; findings that no longer
                                                         # appear are carried forward as "Remediated in PR"
                                                         # when every tool that produced them ran again
        [--dispositions security-assessment/dispositions.json]  # optional Government overrides
        [--fail-on-incomplete]                           # exit 2 when a scanner was skipped/failed or
                                                         # left templates unprocessed (CI gate)
"""
from __future__ import annotations

import argparse
import concurrent.futures
import datetime as dt
import hashlib
import json
import os
import platform
import re
import shutil
import subprocess
import sys
import time
from collections import Counter
from dataclasses import dataclass, field, fields
from pathlib import Path
from typing import Any, Callable, Iterable

import yaml
from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# --------------------------------------------------------------------------------------
# Reference data
# --------------------------------------------------------------------------------------

CAT_I, CAT_II, CAT_III = "CAT I", "CAT II", "CAT III"
SEVERITY_ORDER = {CAT_I: 0, CAT_II: 1, CAT_III: 2}
SEVERITY_BASE_SCORE = {CAT_I: 300, CAT_II: 200, CAT_III: 100}
SEVERITY_JUSTIFICATION = {
    CAT_I: "Direct and immediate loss of confidentiality, integrity or availability (DISA CAT I).",
    CAT_II: "Potential loss of confidentiality, integrity or availability (DISA CAT II).",
    CAT_III: "Degrades protection measures or defense in depth (DISA CAT III).",
}
DISPOSITIONS = ("Open", "Remediated in PR", "Risk acceptance recommended", "Not applicable")

# NIST SP 800-53 Rev. 5 controls used by the rule pack and tool mappings.
NIST_CONTROLS: dict[str, tuple[str, str]] = {
    "AC-2": ("Access Control", "Account Management"),
    "AC-3": ("Access Control", "Access Enforcement"),
    "AC-4": ("Access Control", "Information Flow Enforcement"),
    "AC-6": ("Access Control", "Least Privilege"),
    "AC-6(1)": ("Access Control", "Least Privilege | Authorize Access to Security Functions"),
    "AC-17": ("Access Control", "Remote Access"),
    "AU-2": ("Audit and Accountability", "Event Logging"),
    "AU-3": ("Audit and Accountability", "Content of Audit Records"),
    "AU-6": ("Audit and Accountability", "Audit Record Review, Analysis, and Reporting"),
    "AU-9": ("Audit and Accountability", "Protection of Audit Information"),
    "AU-11": ("Audit and Accountability", "Audit Record Retention"),
    "AU-12": ("Audit and Accountability", "Audit Record Generation"),
    "CM-2": ("Configuration Management", "Baseline Configuration"),
    "CM-3": ("Configuration Management", "Configuration Change Control"),
    "CM-5": ("Configuration Management", "Access Restrictions for Change"),
    "CM-6": ("Configuration Management", "Configuration Settings"),
    "CM-7": ("Configuration Management", "Least Functionality"),
    "CP-9": ("Contingency Planning", "System Backup"),
    "CP-10": ("Contingency Planning", "System Recovery and Reconstitution"),
    "IA-2": ("Identification and Authentication", "Identification and Authentication (Organizational Users)"),
    "IA-5": ("Identification and Authentication", "Authenticator Management"),
    "IA-5(7)": ("Identification and Authentication", "Authenticator Management | No Embedded Unencrypted Static Authenticators"),
    "RA-5": ("Risk Assessment", "Vulnerability Monitoring and Scanning"),
    "SC-5": ("System and Communications Protection", "Denial-of-Service Protection"),
    "SC-7": ("System and Communications Protection", "Boundary Protection"),
    "SC-7(5)": ("System and Communications Protection", "Boundary Protection | Deny by Default — Allow by Exception"),
    "SC-8": ("System and Communications Protection", "Transmission Confidentiality and Integrity"),
    "SC-8(1)": ("System and Communications Protection", "Transmission Confidentiality and Integrity | Cryptographic Protection"),
    "SC-12": ("System and Communications Protection", "Cryptographic Key Establishment and Management"),
    "SC-13": ("System and Communications Protection", "Cryptographic Protection"),
    "SC-28": ("System and Communications Protection", "Protection of Information at Rest"),
    "SC-28(1)": ("System and Communications Protection", "Protection of Information at Rest | Cryptographic Protection"),
    "SI-2": ("System and Information Integrity", "Flaw Remediation"),
    "SI-4": ("System and Information Integrity", "System Monitoring"),
    "SI-7": ("System and Information Integrity", "Software, Firmware, and Information Integrity"),
    "SI-10": ("System and Information Integrity", "Information Input Validation"),
    "SI-12": ("System and Information Integrity", "Information Management and Retention"),
}

# DoD Cloud Computing SRG topic areas (section names from the CC SRG structure) with the matching
# Cloud Computing Mission Owner SRG requirement where one exists.
SRG_AREAS: dict[str, str] = {
    "encryption_rest": "Security Requirements — DoD Policy Regarding Security Controls (FedRAMP+ SC-28); "
                       "Mission Owner SRG SRG-OS-000404-CLD-002720 (encrypt DoD files in cloud storage)",
    "encryption_transit": "Security Requirements — DoD Policy Regarding Security Controls (FedRAMP+ SC-8/SC-13, FIPS 140 validated cryptography)",
    "boundary": "Architecture — Mission Owner network boundary protection (CAP/VDSS); "
                "Mission Owner SRG SRG-OS-000480-CLD-000030 (restrict inbound/outbound traffic flow)",
    "logging": "Computer Network Defense and Incident Response — Continuous Monitoring; "
               "Mission Owner SRG SRG-OS-000342-CLD-000020 (centralized logging)",
    "monitoring": "Computer Network Defense and Incident Response — Continuous Monitoring; "
                  "Mission Owner SRG SRG-NET-000383-CLD-000200 (IDPS / traffic monitoring)",
    "iam": "Identification, Authentication and Access Control; "
           "Mission Owner SRG SRG-OS-000001-CLD-000010 (privileged accounts configured for least privilege)",
    "config": "Security Requirements — Configuration management (CM-6/CM-7); "
              "Mission Owner SRG SRG-OS-000096-CLD-000150 (restrict functions, ports, protocols, services)",
    "recovery": "Data Recovery and Destruction (CP-9 backup and recovery of Mission Owner data)",
    "credentials": "Identification, Authentication and Access Control — authenticator management (IA-5)",
    "validity": "Security Requirements — Configuration management (CM-2/CM-6 baseline configuration)",
}

# CIS AWS Foundations Benchmark v3.0.0 recommendation IDs (verified against the AWS Security Hub CIS
# mapping table: docs.aws.amazon.com/securityhub/latest/userguide/cis-aws-foundations-benchmark.html).
CIS = {
    "s3_tls": "2.1.1",
    "s3_mfa_delete": "2.1.2",
    "s3_block_public": "2.1.4",
    "ebs_encryption": "2.2.1",
    "rds_encryption": "2.3.1",
    "rds_minor_upgrade": "2.3.2",
    "rds_public": "2.3.3",
    "efs_encryption": "2.4.1",
    "cloudtrail_multiregion": "3.1",
    "cloudtrail_validation": "3.2",
    "cloudtrail_bucket_logging": "3.4",
    "cloudtrail_kms": "3.5",
    "kms_rotation": "3.6",
    "vpc_flow_logs": "3.7",
    "nacl_admin_ports": "5.1",
    "sg_admin_ipv4": "5.2",
    "sg_admin_ipv6": "5.3",
    "default_sg": "5.4",
    "imdsv2": "5.6",
}

OWNER_BY_AREA = {
    "encryption_rest": "Cloud Platform Engineering",
    "encryption_transit": "Cloud Platform Engineering",
    "boundary": "Cloud Network Engineering",
    "logging": "Security Operations / ISSO",
    "monitoring": "Security Operations / ISSO",
    "iam": "Identity and Access Management Engineering",
    "config": "Cloud Platform Engineering",
    "recovery": "Cloud Platform Engineering / Data Owner",
    "credentials": "Application Development / ISSO",
    "validity": "Cloud Platform Engineering",
}
TARGET_DATE_PLACEHOLDER = {
    CAT_I: "TBD — CAT I: 30 days from Government acceptance of this report",
    CAT_II: "TBD — CAT II: 90 days from Government acceptance of this report",
    CAT_III: "TBD — CAT III: 180 days from Government acceptance of this report",
}

ADMIN_PORTS = {22, 3389, 5985, 5986}
DATABASE_PORTS = {1433, 1521, 3306, 5432, 5439, 6379, 8020, 9200, 9300, 11211, 27017, 27018}
WEB_PORTS = {80, 443, 8080, 8443}
PUBLIC_CIDRS = {"0.0.0.0/0", "::/0"}

TEMPLATE_EXTS = {".yaml", ".yml", ".json", ".template"}
SKIP_DIRS = {".git", "node_modules", "security-assessment", "cdk.out", ".venv", "venv", "__pycache__"}


class _Unknown:
    def __repr__(self) -> str:
        return "<unresolved>"


UNKNOWN = _Unknown()

# --------------------------------------------------------------------------------------
# Line-aware CloudFormation YAML loading
# --------------------------------------------------------------------------------------


class LMap(dict):
    """dict that remembers the source line of the mapping and of each key."""

    __line__: int = 0
    __key_lines__: dict

    def line_of(self, key: str) -> int | None:
        return getattr(self, "__key_lines__", {}).get(key)


class LList(list):
    __line__: int = 0


class CfnLoader(yaml.SafeLoader):
    pass


def _construct_mapping(loader: CfnLoader, node: yaml.MappingNode, deep: bool = False) -> LMap:  # pylint: disable=unused-argument
    loader.flatten_mapping(node)
    m = LMap()
    m.__line__ = node.start_mark.line + 1
    m.__key_lines__ = {}
    for key_node, value_node in node.value:
        key = loader.construct_object(key_node, deep=True)
        value = loader.construct_object(value_node, deep=True)
        if isinstance(key, (list, dict)):
            key = str(key)
        m[key] = value
        m.__key_lines__[key] = key_node.start_mark.line + 1
    return m


def _construct_sequence(loader: CfnLoader, node: yaml.SequenceNode) -> LList:
    lst = LList(loader.construct_object(child, deep=True) for child in node.value)
    lst.__line__ = node.start_mark.line + 1
    return lst


def _construct_tag(loader: CfnLoader, tag_suffix: str, node: yaml.Node) -> Any:
    if isinstance(node, yaml.ScalarNode):
        value: Any = loader.construct_scalar(node)
    elif isinstance(node, yaml.SequenceNode):
        value = _construct_sequence(loader, node)
    else:
        value = _construct_mapping(loader, node)
    if tag_suffix == "Ref":
        key = "Ref"
    elif tag_suffix == "Condition":
        key = "Condition"
    elif tag_suffix == "GetAtt":
        key = "Fn::GetAtt"
        if isinstance(value, str):
            value = value.split(".", 1)
    elif tag_suffix.startswith("Rain::"):
        key = tag_suffix
    else:
        key = "Fn::" + tag_suffix
    m = LMap({key: value})
    m.__line__ = node.start_mark.line + 1
    m.__key_lines__ = {key: node.start_mark.line + 1}
    return m


CfnLoader.add_constructor(yaml.resolver.BaseResolver.DEFAULT_MAPPING_TAG, _construct_mapping)
CfnLoader.add_constructor(yaml.resolver.BaseResolver.DEFAULT_SEQUENCE_TAG, _construct_sequence)
CfnLoader.add_multi_constructor("!", _construct_tag)
CfnLoader.add_constructor("tag:yaml.org,2002:timestamp", lambda loader, node: loader.construct_scalar(node))
# CloudFormation only treats true/false as booleans; yes/no/on/off stay strings (YAML 1.2 behaviour).
CfnLoader.yaml_implicit_resolvers = {
    k: [(tag, rx) for tag, rx in v if tag != "tag:yaml.org,2002:bool"] for k, v in CfnLoader.yaml_implicit_resolvers.items()
}
CfnLoader.add_implicit_resolver("tag:yaml.org,2002:bool", re.compile(r"^(?:true|True|TRUE|false|False|FALSE)$"), list("tTfF"))


class TemplateParseError(ValueError):
    """The file is not well-formed JSON or YAML (as opposed to well-formed content that is not a template)."""


def parse_documents(path: Path) -> list[Any]:
    """Every document in the file; raises TemplateParseError when the text is malformed."""
    text = path.read_text(encoding="utf-8", errors="replace")
    if path.suffix == ".json" or text.lstrip().startswith("{"):
        try:
            return [json.loads(text)]
        except json.JSONDecodeError as e:
            raise TemplateParseError(str(e)) from e
    try:
        return list(yaml.load_all(text, Loader=CfnLoader))  # noqa: S506 - SafeLoader subclass
    except yaml.YAMLError as e:
        raise TemplateParseError(str(e)) from e


def load_template(path: Path) -> dict | None:
    """The single mapping document in the file, or None when the file is malformed or is not one mapping."""
    try:
        docs = parse_documents(path)
    except TemplateParseError:
        return None
    return docs[0] if len(docs) == 1 and isinstance(docs[0], dict) else None


FOREACH_PREFIX = "Fn::ForEach::"


def expand_resources(resources: Any) -> dict[str, dict]:
    """The Resources section with ``AWS::LanguageExtensions`` ``Fn::ForEach`` loops expanded.

    ``Fn::ForEach::<Name>: [Identifier, Collection, {"Prefix${Identifier}Suffix": <resource>}]`` yields one resource
    per collection item with ``${Identifier}`` substituted in the logical ID. When the collection is not a literal
    list (for example a ``!Ref`` to a CommaDelimitedList parameter) the fragment is kept once with the ``${Identifier}``
    placeholder in its logical ID, so its properties are still assessed. Loops nest."""
    out: dict[str, dict] = {}
    if not isinstance(resources, dict):
        return out
    for key, value in resources.items():
        if not (isinstance(key, str) and key.startswith(FOREACH_PREFIX)):
            if isinstance(value, dict):
                out[key] = value
            continue
        if not (isinstance(value, list) and len(value) == 3 and isinstance(value[0], str) and isinstance(value[2], dict)):
            continue
        identifier, collection, fragment = value
        items = [str(i) for i in collection] if isinstance(collection, list) and all(isinstance(i, (str, int)) for i in collection) else [f"${{{identifier}}}"]
        placeholder = f"${{{identifier}}}"
        for item in items:
            for lid, res in expand_resources(fragment).items():
                out[lid.replace(placeholder, item)] = res
    return out


def is_cfn_template(data: dict) -> bool:
    return any("Type" in r for r in expand_resources(data.get("Resources")).values())


# --------------------------------------------------------------------------------------
# Template context and helpers
# --------------------------------------------------------------------------------------


_UNSET = object()


@dataclass
class Hit:
    logical_id: str
    resource_type: str
    line: int
    detail: str
    severity: str | None = None  # overrides the rule default
    disposition: str | None = None
    title: str | None = None
    cis: str | None | object = _UNSET  # overrides the rule CIS ID (None clears it)
    key: str | None = None  # identity of this hit within the resource for rules that can report one resource several times


@dataclass
class TemplateCtx:
    path: Path
    rel: str
    text: str
    lines: list[str]
    data: dict
    service_dir: str

    @property
    def resources(self) -> dict[str, dict]:
        return expand_resources(self.data.get("Resources"))

    @property
    def parameters(self) -> dict[str, dict]:
        p = self.data.get("Parameters") or {}
        return {k: v for k, v in p.items() if isinstance(v, dict)} if isinstance(p, dict) else {}

    def by_type(self, *types: str) -> list[tuple[str, dict]]:
        return [(k, v) for k, v in self.resources.items() if v.get("Type") in types]

    def has_type(self, *types: str) -> bool:
        return bool(self.by_type(*types))

    def find_line(self, pattern: str, start: int = 0) -> int | None:
        rx = re.compile(pattern)
        for i in range(max(start, 0), len(self.lines)):
            if rx.search(self.lines[i]):
                return i + 1
        return None

    def resource_line(self, logical_id: str, res: dict) -> int:
        if isinstance(res, LMap) and res.__line__:
            return res.__line__
        return self.find_line(rf'^\s*"?{re.escape(logical_id)}"?\s*:') or 1

    def prop_line(self, logical_id: str, res: dict, *keys: str) -> int:
        """Line of the deepest existing key along Properties.<keys>, else the resource line."""
        node: Any = res
        line = self.resource_line(logical_id, res)
        for key in keys:
            if isinstance(node, LMap):
                kl = node.line_of(key)
                if kl:
                    line = kl
            if isinstance(node, dict) and key in node:
                node = node[key]
            else:
                break
        if not isinstance(res, LMap) and keys:  # JSON: search textually after the resource line
            found = self.find_line(rf'"{re.escape(str(keys[-1]))}"\s*:', start=line - 1)
            if found:
                line = found
        return line

    def snippet(self, line: int, width: int = 140) -> str:
        if 1 <= line <= len(self.lines):
            s = self.lines[line - 1].strip()
            return s if len(s) <= width else s[: width - 3] + "..."
        return ""

    def resolve(self, value: Any, depth: int = 0) -> list[Any]:
        """Possible literal values of a property (parameter defaults, Fn::If branches)."""
        if depth > 6:
            return [UNKNOWN]
        if isinstance(value, dict):
            if "Ref" in value and isinstance(value["Ref"], str):
                param = self.parameters.get(value["Ref"])
                if param is None:
                    return [UNKNOWN]
                return [param["Default"]] if "Default" in param else [UNKNOWN]
            if "Fn::If" in value and isinstance(value["Fn::If"], list) and len(value["Fn::If"]) == 3:
                out: list[Any] = []
                for branch in value["Fn::If"][1:]:
                    out.extend(self.resolve(branch, depth + 1))
                return out
            return [UNKNOWN]
        return [value]

    def param_for(self, value: Any) -> str | None:
        if isinstance(value, dict) and isinstance(value.get("Ref"), str) and value["Ref"] in self.parameters:
            return value["Ref"]
        return None


def props(res: dict) -> dict:
    p = res.get("Properties")
    return p if isinstance(p, dict) else {}


def rtype_of(res: dict) -> str:
    t = res.get("Type", "")
    return t if isinstance(t, str) else json.dumps(t, default=str)


def is_true(v: Any) -> bool:
    return v is True or (isinstance(v, str) and v.strip().lower() == "true")


def any_true(ctx: TemplateCtx, v: Any) -> bool:
    return any(is_true(x) for x in ctx.resolve(v))


def all_true(ctx: TemplateCtx, v: Any) -> bool:
    vals = ctx.resolve(v)
    return bool(vals) and all(is_true(x) for x in vals)


def literal_str(v: Any) -> str | None:
    return v if isinstance(v, str) else None


def as_list(v: Any) -> list:
    if v is None:
        return []
    return list(v) if isinstance(v, list) else [v]


def iter_statements(policy_doc: Any) -> Iterable[dict]:
    if not isinstance(policy_doc, dict):
        return
    for s in as_list(policy_doc.get("Statement")):
        if isinstance(s, dict):
            yield s


def stmt_actions(s: dict) -> list[str]:
    return [a for a in as_list(s.get("Action")) if isinstance(a, str)]


def stmt_resources(s: dict) -> list[Any]:
    return as_list(s.get("Resource"))


def principal_is_public(p: Any) -> bool:
    if p == "*":
        return True
    if isinstance(p, dict):
        aws = p.get("AWS")
        return aws == "*" or (isinstance(aws, list) and "*" in aws)
    return False


def statement_has_condition(s: dict) -> bool:
    c = s.get("Condition")
    return isinstance(c, dict) and len(c) > 0


def node_line(node: Any, fallback: int) -> int:
    return node.__line__ if isinstance(node, (LMap, LList)) and node.__line__ else fallback


def parameter_looks_secret(name: str, param: dict) -> bool:
    """True when a parameter carries an authenticator value (not a name, ARN or flag that merely mentions one)."""
    if param.get("AllowedValues") or str(param.get("Type", "String")) not in {"String", "AWS::SSM::Parameter::Value<String>"}:
        return False
    lname = name.lower()
    if re.search(r"(name|arn|id|path|prefix|bucket|url|uri|endpoint|enabled?|required|create|use|enable)$", lname):
        return False
    if re.search(r"keyname|key ?pair|kmskey|kms|publickey|public key|bucketkey|apikeyrequired|keyid|key id|s3key|objectkey|keyarn", lname):
        return False
    if re.search(r"passw|secret|token|apikey|credential|privatekey", lname):
        return True
    desc = str(param.get("Description", "")).lower()
    return bool(re.search(r"\bpassword\b|\bpassphrase\b|\bsecret key\b|\baccess key\b|\bprivate key\b", desc)) and not re.search(r"\b(name|arn) of\b", desc)


# --------------------------------------------------------------------------------------
# Rule pack
# --------------------------------------------------------------------------------------


@dataclass
class Rule:
    id: str
    title: str
    severity: str
    nist: list[str]
    area: str
    description: str
    remediation: str
    applies_to: list[str]
    cis: str | None = None
    overlaps: set[str] = field(default_factory=set)  # Checkov / cfn_nag IDs that report the same weakness
    tags: set[str] = field(default_factory=set)  # exposure, credential, datastore
    check: Callable[[TemplateCtx], Iterable[Hit]] | None = None


RULES: list[Rule] = []


def rule(**kwargs):
    def deco(fn):
        RULES.append(Rule(check=fn, **kwargs))
        return fn

    return deco


# ---- Encryption at rest -------------------------------------------------------------------

S3_BUCKET = "AWS::S3::Bucket"


@rule(id="CSA-ENC-001", title="S3 bucket does not declare server-side encryption",
      severity=CAT_II, nist=["SC-28", "SC-28(1)"], area="encryption_rest",
      applies_to=[S3_BUCKET], overlaps={"CKV_AWS_19", "W41"}, tags={"datastore"},
      description="The bucket has no BucketEncryption block. Encryption then depends on the account default (SSE-S3) and the baseline cannot show that data at rest is protected with an approved key.",
      remediation="Add BucketEncryption with ServerSideEncryptionByDefault using aws:kms and a customer-managed key; set BucketKeyEnabled: true.")
def r_s3_encryption(ctx: TemplateCtx):
    for lid, res in ctx.by_type(S3_BUCKET):
        if "BucketEncryption" not in props(res):
            yield Hit(lid, S3_BUCKET, ctx.resource_line(lid, res), "Properties.BucketEncryption is absent")


@rule(id="CSA-ENC-002", title="S3 bucket encryption does not use a customer-managed KMS key",
      severity=CAT_III, nist=["SC-28(1)", "SC-12"], area="encryption_rest",
      applies_to=[S3_BUCKET], overlaps={"CKV_AWS_145"}, tags={"datastore"},
      description="Server-side encryption is enabled but uses SSE-S3 or an AWS-managed key. The Government cannot control key policy, rotation or revocation.",
      remediation="Set SSEAlgorithm: aws:kms and KMSMasterKeyID to a customer-managed key with a restrictive key policy.")
def r_s3_cmk(ctx: TemplateCtx):
    for lid, res in ctx.by_type(S3_BUCKET):
        enc = props(res).get("BucketEncryption")
        if not isinstance(enc, dict):
            continue
        for rule_ in as_list(enc.get("ServerSideEncryptionConfiguration")):
            if not isinstance(rule_, dict):
                continue
            d = rule_.get("ServerSideEncryptionByDefault")
            if isinstance(d, dict) and (d.get("SSEAlgorithm") == "AES256" or not d.get("KMSMasterKeyID")):
                yield Hit(lid, S3_BUCKET, ctx.prop_line(lid, res, "Properties", "BucketEncryption"),
                          f"SSEAlgorithm={d.get('SSEAlgorithm')} without KMSMasterKeyID")
                break


DB_TYPES = {
    "AWS::RDS::DBInstance": "StorageEncrypted",
    "AWS::RDS::DBCluster": "StorageEncrypted",
    "AWS::Redshift::Cluster": "Encrypted",
    "AWS::DocDB::DBCluster": "StorageEncrypted",
    "AWS::Neptune::DBCluster": "StorageEncrypted",
}


@rule(id="CSA-ENC-003", title="Database storage is not encrypted at rest",
      severity=CAT_I, nist=["SC-28", "SC-28(1)"], area="encryption_rest", cis=CIS["rds_encryption"],
      applies_to=list(DB_TYPES), overlaps={"CKV_AWS_16", "CKV_AWS_96", "CKV_AWS_64", "CKV_AWS_74", "CKV_AWS_44", "F27", "F26", "F28"},
      tags={"datastore"},
      description="The database resource does not set storage encryption to true. Mission data, automated backups and snapshots are stored in clear text. The Mission Owner SRG rates encryption of DoD data in cloud storage as High.",
      remediation="Set StorageEncrypted: true (Encrypted: true for Redshift) and supply KmsKeyId with a customer-managed key. Encryption is set at creation; existing instances need a snapshot-copy migration.")
def r_db_encryption(ctx: TemplateCtx):
    for lid, res in ctx.resources.items():
        rtype = rtype_of(res)
        if rtype not in DB_TYPES:
            continue
        p = props(res)
        if rtype == "AWS::RDS::DBInstance":
            if p.get("SourceDBInstanceIdentifier") or p.get("DBClusterIdentifier") or p.get("DBSnapshotIdentifier"):
                continue  # inherits from source / cluster / snapshot
            if (literal_str(p.get("Engine")) or "").startswith("aurora"):
                continue
        if p.get("SnapshotIdentifier") or p.get("SourceDBClusterIdentifier"):
            continue
        key = DB_TYPES[rtype]
        if not all_true(ctx, p.get(key)):
            detail = f"Properties.{key} is absent" if key not in p else f"Properties.{key} resolves to {ctx.resolve(p.get(key))}"
            # CIS 2.3.1 covers RDS only; other engines carry no CIS recommendation ID.
            yield Hit(lid, rtype, ctx.prop_line(lid, res, "Properties", key), detail,
                      cis=CIS["rds_encryption"] if rtype.startswith("AWS::RDS::") else None)


EBS_HOSTS = ["AWS::EC2::Instance", "AWS::EC2::LaunchTemplate", "AWS::AutoScaling::LaunchConfiguration"]


@rule(id="CSA-ENC-004", title="EBS volume is not encrypted",
      severity=CAT_II, nist=["SC-28", "SC-28(1)"], area="encryption_rest", cis=CIS["ebs_encryption"],
      applies_to=["AWS::EC2::Volume"] + EBS_HOSTS, overlaps={"CKV_AWS_3", "CKV_AWS_8", "CKV_AWS_189", "F1"}, tags={"datastore"},
      description="An EBS volume or block device mapping does not set Encrypted: true. Data on the volume and its snapshots is unprotected unless account-level default encryption is enabled.",
      remediation="Set Encrypted: true (and KmsKeyId) on every AWS::EC2::Volume and on each BlockDeviceMappings[].Ebs entry; enable EBS default encryption in the account as a backstop.")
def r_ebs_encryption(ctx: TemplateCtx):
    for lid, res in ctx.by_type("AWS::EC2::Volume"):
        p = props(res)
        if not all_true(ctx, p.get("Encrypted")):
            yield Hit(lid, "AWS::EC2::Volume", ctx.prop_line(lid, res, "Properties", "Encrypted"),
                      "Properties.Encrypted is absent" if "Encrypted" not in p else f"Encrypted resolves to {ctx.resolve(p.get('Encrypted'))}")
    for lid, res in ctx.resources.items():
        rtype = rtype_of(res)
        if rtype not in EBS_HOSTS:
            continue
        p = props(res)
        container = p.get("LaunchTemplateData") if rtype == "AWS::EC2::LaunchTemplate" else p
        if not isinstance(container, dict):
            continue
        for i, bdm in enumerate(as_list(container.get("BlockDeviceMappings"))):
            if not isinstance(bdm, dict):
                continue
            ebs = bdm.get("Ebs")
            if isinstance(ebs, dict) and not all_true(ctx, ebs.get("Encrypted")):
                yield Hit(lid, rtype, node_line(ebs, ctx.prop_line(lid, res, "Properties", "BlockDeviceMappings")),
                          f"BlockDeviceMappings[{i}].Ebs.Encrypted is not true")
                break


@rule(id="CSA-ENC-005", title="EFS file system is not encrypted at rest",
      severity=CAT_II, nist=["SC-28", "SC-28(1)"], area="encryption_rest", cis=CIS["efs_encryption"],
      applies_to=["AWS::EFS::FileSystem"], overlaps={"CKV_AWS_42", "CKV_AWS_184"}, tags={"datastore"},
      description="The EFS file system does not set Encrypted: true. File data is stored in clear text and encryption cannot be enabled after creation.",
      remediation="Set Encrypted: true and KmsKeyId to a customer-managed key.")
def r_efs_encryption(ctx: TemplateCtx):
    for lid, res in ctx.by_type("AWS::EFS::FileSystem"):
        if not all_true(ctx, props(res).get("Encrypted")):
            yield Hit(lid, "AWS::EFS::FileSystem", ctx.prop_line(lid, res, "Properties", "Encrypted"), "Properties.Encrypted is not true")


@rule(id="CSA-ENC-006", title="SQS queue does not declare server-side encryption",
      severity=CAT_III, nist=["SC-28", "SC-28(1)"], area="encryption_rest",
      applies_to=["AWS::SQS::Queue"], overlaps={"CKV_AWS_27", "W48"}, tags={"datastore"},
      description="The queue sets neither KmsMasterKeyId nor SqsManagedSseEnabled. Protection depends on the service default (SSE-SQS) and the Government cannot control the key.",
      remediation="Set KmsMasterKeyId to a customer-managed key (or at minimum SqsManagedSseEnabled: true).")
def r_sqs_encryption(ctx: TemplateCtx):
    for lid, res in ctx.by_type("AWS::SQS::Queue"):
        p = props(res)
        if not p.get("KmsMasterKeyId") and not any_true(ctx, p.get("SqsManagedSseEnabled")):
            yield Hit(lid, "AWS::SQS::Queue", ctx.resource_line(lid, res), "KmsMasterKeyId and SqsManagedSseEnabled are absent")


@rule(id="CSA-ENC-007", title="SNS topic is not encrypted at rest",
      severity=CAT_II, nist=["SC-28", "SC-28(1)"], area="encryption_rest",
      applies_to=["AWS::SNS::Topic"], overlaps={"CKV_AWS_26", "W47"}, tags={"datastore"},
      description="The topic has no KmsMasterKeyId. Message payloads are stored unencrypted by the service.",
      remediation="Set KmsMasterKeyId to a customer-managed key and grant the publishing services kms:GenerateDataKey/kms:Decrypt in the key policy.")
def r_sns_encryption(ctx: TemplateCtx):
    for lid, res in ctx.by_type("AWS::SNS::Topic"):
        if not props(res).get("KmsMasterKeyId"):
            yield Hit(lid, "AWS::SNS::Topic", ctx.resource_line(lid, res), "Properties.KmsMasterKeyId is absent")


@rule(id="CSA-ENC-008", title="CloudWatch log group is not encrypted with a customer-managed key",
      severity=CAT_III, nist=["SC-28(1)", "AU-9"], area="encryption_rest",
      applies_to=["AWS::Logs::LogGroup"], overlaps={"CKV_AWS_158", "W84"},
      description="The log group has no KmsKeyId. Log data is protected only by the service-owned key, and audit records may contain sensitive content.",
      remediation="Set KmsKeyId to a customer-managed key whose policy grants the logs service principal for the Region.")
def r_loggroup_kms(ctx: TemplateCtx):
    for lid, res in ctx.by_type("AWS::Logs::LogGroup"):
        if not props(res).get("KmsKeyId"):
            yield Hit(lid, "AWS::Logs::LogGroup", ctx.resource_line(lid, res), "Properties.KmsKeyId is absent")


@rule(id="CSA-ENC-009", title="Streaming resource is not encrypted at rest",
      severity=CAT_II, nist=["SC-28", "SC-28(1)"], area="encryption_rest",
      applies_to=["AWS::Kinesis::Stream", "AWS::KinesisFirehose::DeliveryStream"], overlaps={"CKV_AWS_43", "CKV_AWS_241", "CKV_AWS_240"},
      tags={"datastore"},
      description="The Kinesis stream has no StreamEncryption, or the Firehose delivery stream (DirectPut) has no DeliveryStreamEncryptionConfigurationInput. Buffered records are stored unencrypted.",
      remediation="Add StreamEncryption (EncryptionType: KMS, KeyId) or DeliveryStreamEncryptionConfigurationInput (KeyType: CUSTOMER_MANAGED_CMK).")
def r_stream_encryption(ctx: TemplateCtx):
    for lid, res in ctx.by_type("AWS::Kinesis::Stream"):
        if "StreamEncryption" not in props(res):
            yield Hit(lid, "AWS::Kinesis::Stream", ctx.resource_line(lid, res), "Properties.StreamEncryption is absent")
    for lid, res in ctx.by_type("AWS::KinesisFirehose::DeliveryStream"):
        p = props(res)
        if p.get("DeliveryStreamType", "DirectPut") == "DirectPut" and "DeliveryStreamEncryptionConfigurationInput" not in p:
            yield Hit(lid, "AWS::KinesisFirehose::DeliveryStream", ctx.resource_line(lid, res), "DeliveryStreamEncryptionConfigurationInput is absent")


@rule(id="CSA-ENC-010", title="DynamoDB table does not use a customer-managed KMS key",
      severity=CAT_III, nist=["SC-28(1)", "SC-12"], area="encryption_rest",
      applies_to=["AWS::DynamoDB::Table"], overlaps={"CKV_AWS_119", "W74"}, tags={"datastore"},
      description="The table has no SSESpecification with SSEType KMS and a KMSMasterKeyId. DynamoDB encrypts with an AWS-owned key by default, which the Government cannot audit or revoke.",
      remediation="Add SSESpecification: {SSEEnabled: true, SSEType: KMS, KMSMasterKeyId: <customer-managed key>}.")
def r_dynamodb_cmk(ctx: TemplateCtx):
    for lid, res in ctx.by_type("AWS::DynamoDB::Table"):
        sse = props(res).get("SSESpecification")
        if not (isinstance(sse, dict) and any_true(ctx, sse.get("SSEEnabled")) and sse.get("KMSMasterKeyId")):
            yield Hit(lid, "AWS::DynamoDB::Table", ctx.resource_line(lid, res), "SSESpecification with a customer-managed key is absent")


@rule(id="CSA-ENC-011", title="CloudTrail trail logs are not encrypted with KMS",
      severity=CAT_II, nist=["AU-9", "SC-28(1)"], area="encryption_rest", cis=CIS["cloudtrail_kms"],
      applies_to=["AWS::CloudTrail::Trail"], overlaps={"CKV_AWS_35"},
      description="The trail has no KMSKeyId. Audit logs are protected only by S3 default encryption and the Government cannot restrict who can decrypt them.",
      remediation="Set KMSKeyId to a customer-managed key with a policy that allows cloudtrail.amazonaws.com to encrypt.")
def r_cloudtrail_kms(ctx: TemplateCtx):
    for lid, res in ctx.by_type("AWS::CloudTrail::Trail"):
        if not props(res).get("KMSKeyId"):
            yield Hit(lid, "AWS::CloudTrail::Trail", ctx.resource_line(lid, res), "Properties.KMSKeyId is absent")


@rule(id="CSA-ENC-012", title="Cache or search domain is not encrypted at rest",
      severity=CAT_II, nist=["SC-28", "SC-28(1)"], area="encryption_rest",
      applies_to=["AWS::ElastiCache::ReplicationGroup", "AWS::Elasticsearch::Domain", "AWS::OpenSearchService::Domain"],
      overlaps={"CKV_AWS_29", "CKV_AWS_5", "CKV_AWS_247", "CKV_AWS_31", "F25", "F33"}, tags={"datastore"},
      description="AtRestEncryptionEnabled (ElastiCache) or EncryptionAtRestOptions.Enabled (OpenSearch) is not true.",
      remediation="Set AtRestEncryptionEnabled: true / EncryptionAtRestOptions: {Enabled: true, KmsKeyId}.")
def r_cache_search_rest(ctx: TemplateCtx):
    for lid, res in ctx.by_type("AWS::ElastiCache::ReplicationGroup"):
        if not all_true(ctx, props(res).get("AtRestEncryptionEnabled")):
            yield Hit(lid, rtype_of(res), ctx.prop_line(lid, res, "Properties", "AtRestEncryptionEnabled"), "AtRestEncryptionEnabled is not true")
    for lid, res in ctx.by_type("AWS::Elasticsearch::Domain", "AWS::OpenSearchService::Domain"):
        o = props(res).get("EncryptionAtRestOptions")
        if not (isinstance(o, dict) and all_true(ctx, o.get("Enabled"))):
            yield Hit(lid, rtype_of(res), ctx.resource_line(lid, res), "EncryptionAtRestOptions.Enabled is not true")


@rule(id="CSA-KMS-001", title="Encrypted resource relies on an AWS-managed key instead of a customer-managed key",
      severity=CAT_III, nist=["SC-12", "SC-28(1)"], area="encryption_rest",
      applies_to=["AWS::RDS::DBInstance", "AWS::RDS::DBCluster", "AWS::EFS::FileSystem", "AWS::EC2::Volume", "AWS::Redshift::Cluster"],
      description="Encryption is enabled but no KmsKeyId is set, so the service default AWS-managed key is used. Key policy, rotation and cross-account use cannot be controlled.",
      remediation="Set KmsKeyId to a customer-managed key created in the template or supplied as a parameter.")
def r_aws_managed_key(ctx: TemplateCtx):
    flags = {"AWS::RDS::DBInstance": "StorageEncrypted", "AWS::RDS::DBCluster": "StorageEncrypted", "AWS::EFS::FileSystem": "Encrypted",
             "AWS::EC2::Volume": "Encrypted", "AWS::Redshift::Cluster": "Encrypted"}
    for lid, res in ctx.resources.items():
        flag = flags.get(rtype_of(res))
        p = props(res)
        if flag and all_true(ctx, p.get(flag)) and not p.get("KmsKeyId"):
            yield Hit(lid, rtype_of(res), ctx.prop_line(lid, res, "Properties", flag), f"{flag} is true but KmsKeyId is absent")


@rule(id="CSA-KMS-002", title="KMS key does not enable automatic rotation",
      severity=CAT_III, nist=["SC-12"], area="encryption_rest", cis=CIS["kms_rotation"],
      applies_to=["AWS::KMS::Key"], overlaps={"CKV_AWS_7"},
      description="EnableKeyRotation is not true on a symmetric customer-managed key.",
      remediation="Set EnableKeyRotation: true.")
def r_kms_rotation(ctx: TemplateCtx):
    for lid, res in ctx.by_type("AWS::KMS::Key"):
        p = props(res)
        if p.get("KeySpec", "SYMMETRIC_DEFAULT") != "SYMMETRIC_DEFAULT":
            continue
        if not all_true(ctx, p.get("EnableKeyRotation")):
            yield Hit(lid, "AWS::KMS::Key", ctx.resource_line(lid, res), "EnableKeyRotation is not true")


# ---- Encryption in transit ----------------------------------------------------------------


def bucket_policy_targets(ctx: TemplateCtx, bucket_ref: Any) -> set[str]:
    """Logical IDs of the buckets a BucketPolicy.Bucket expression can denote.

    ``!Ref Bucket`` names the bucket directly. A literal or ``!Sub`` name matches every bucket whose ``BucketName``
    is the same expression (after canonicalisation), so an explicitly named bucket and its policy are associated
    the way CloudFormation associates them at deploy time. An expression that matches several buckets is
    ambiguous and matches none."""
    if isinstance(bucket_ref, dict) and isinstance(bucket_ref.get("Ref"), str):
        lid = bucket_ref["Ref"]
        return {lid} if lid in ctx.resources else set()
    wanted = _canon(bucket_ref)
    matches = {lid for lid, res in ctx.by_type(S3_BUCKET) if "BucketName" in props(res) and _canon(props(res)["BucketName"]) == wanted}
    return matches if len(matches) == 1 else set()


def bucket_policies_for(ctx: TemplateCtx, bucket_lid: str) -> list[tuple[str, dict]]:
    return [(plid, pres) for plid, pres in ctx.by_type("AWS::S3::BucketPolicy")
            if bucket_lid in bucket_policy_targets(ctx, props(pres).get("Bucket"))]


def policy_denies_insecure_transport(doc: Any) -> bool:
    """True only for a bucket-wide deny: Principal *, Action s3:* (or *), bucket and object ARNs, aws:SecureTransport=false.

    Statement scopes are combined, so two deny statements that split bucket and object ARNs also qualify.
    """
    bucket = objects = False
    for s in iter_statements(doc):
        if s.get("Effect") != "Deny" or not principal_is_public(s.get("Principal")):
            continue
        cond = s.get("Condition")
        b = cond.get("Bool") if isinstance(cond, dict) else None
        if not (isinstance(b, dict) and str(b.get("aws:SecureTransport", "")).lower() == "false"):
            continue
        if len(cond) > 1:
            continue  # further conditions narrow the deny to a subset of requests
        if not any(a in {"*", "s3:*"} for a in stmt_actions(s)):
            continue
        for r in stmt_resources(s):
            rendered = r if isinstance(r, str) else json.dumps(r, default=str)
            if rendered.strip() == "*":
                return True
            if "/*" in rendered:
                objects = True
            else:
                bucket = True
    return bucket and objects


@rule(id="CSA-TLS-001", title="S3 bucket does not deny non-TLS (aws:SecureTransport=false) requests",
      severity=CAT_II, nist=["SC-8", "SC-8(1)"], area="encryption_transit", cis=CIS["s3_tls"],
      applies_to=[S3_BUCKET], tags={"datastore"},
      description="No bucket policy in the template denies requests when aws:SecureTransport is false. Clients may read or write objects over clear-text HTTP.",
      remediation="Attach an AWS::S3::BucketPolicy with a Deny statement for s3:* on the bucket and its objects with Condition Bool aws:SecureTransport: false.")
def r_s3_tls(ctx: TemplateCtx):
    for lid, res in ctx.by_type(S3_BUCKET):
        if any(policy_denies_insecure_transport(props(pres).get("PolicyDocument")) for _, pres in bucket_policies_for(ctx, lid)):
            continue
        yield Hit(lid, S3_BUCKET, ctx.resource_line(lid, res), "no BucketPolicy denies aws:SecureTransport=false for this bucket")


def listener_redirects_to_https(p: dict) -> bool:
    for a in as_list(p.get("DefaultActions")):
        if isinstance(a, dict) and a.get("Type") == "redirect":
            rc = a.get("RedirectConfig") or {}
            if isinstance(rc, dict) and rc.get("Protocol") == "HTTPS":
                return True
    return False


@rule(id="CSA-TLS-002", title="Load balancer listener accepts clear-text traffic",
      severity=CAT_II, nist=["SC-8", "SC-8(1)"], area="encryption_transit",
      applies_to=["AWS::ElasticLoadBalancingV2::Listener", "AWS::ElasticLoadBalancing::LoadBalancer"],
      overlaps={"CKV_AWS_2", "W56"}, tags={"exposure"},
      description="The listener uses HTTP (or a TCP/HTTP Classic ELB listener) and does not redirect to HTTPS. Data between clients and the load balancer is not encrypted.",
      remediation="Use Protocol: HTTPS/TLS with an ACM certificate and SslPolicy ELBSecurityPolicy-TLS13-1-2-2021-06; keep a port-80 listener only as a redirect to HTTPS.")
def r_lb_cleartext(ctx: TemplateCtx):
    for lid, res in ctx.by_type("AWS::ElasticLoadBalancingV2::Listener"):
        p = props(res)
        proto = literal_str(p.get("Protocol")) or ""
        if proto.upper() in {"HTTP", "TCP", "UDP", "TCP_UDP"} and not listener_redirects_to_https(p):
            yield Hit(lid, rtype_of(res), ctx.prop_line(lid, res, "Properties", "Protocol"), f"Protocol={proto} without HTTPS redirect")
    for lid, res in ctx.by_type("AWS::ElasticLoadBalancing::LoadBalancer"):
        for i, lst in enumerate(as_list(props(res).get("Listeners"))):
            if isinstance(lst, dict) and str(lst.get("Protocol", "")).upper() in {"HTTP", "TCP"}:
                yield Hit(lid, rtype_of(res), node_line(lst, ctx.prop_line(lid, res, "Properties", "Listeners")),
                          f"Listeners[{i}].Protocol={lst.get('Protocol')} (clear text)")
                break


STRONG_TLS_POLICY = re.compile(r"^ELBSecurityPolicy-(TLS13-1-2|TLS13-1-3|TLS-1-2|FS-1-2)")


@rule(id="CSA-TLS-003", title="HTTPS listener does not enforce a TLS 1.2+ security policy",
      severity=CAT_II, nist=["SC-8(1)", "SC-13"], area="encryption_transit",
      applies_to=["AWS::ElasticLoadBalancingV2::Listener"], overlaps={"CKV_AWS_103", "W55"},
      description="The HTTPS/TLS listener has no SslPolicy or uses a policy that still permits TLS 1.0/1.1 (for example the default ELBSecurityPolicy-2016-08).",
      remediation="Set SslPolicy: ELBSecurityPolicy-TLS13-1-2-2021-06 (or another TLS 1.2-minimum policy).")
def r_lb_tls_policy(ctx: TemplateCtx):
    for lid, res in ctx.by_type("AWS::ElasticLoadBalancingV2::Listener"):
        p = props(res)
        if (literal_str(p.get("Protocol")) or "").upper() not in {"HTTPS", "TLS"}:
            continue
        pol = p.get("SslPolicy")
        for v in (ctx.resolve(pol) if pol is not None else [None]):
            if v is UNKNOWN:
                continue
            if v is None or not STRONG_TLS_POLICY.match(str(v)):
                yield Hit(lid, rtype_of(res), ctx.prop_line(lid, res, "Properties", "SslPolicy"), f"SslPolicy={v}")
                break


FORCE_SSL_PARAMS = {"rds.force_ssl", "require_secure_transport"}


def param_group_forces_ssl(pg: dict) -> bool:
    params = props(pg).get("Parameters")
    if not isinstance(params, dict):
        return False
    return any(k in FORCE_SSL_PARAMS and str(v).strip().lower() in {"1", "on", "true"} for k, v in params.items())


@rule(id="CSA-TLS-004", title="RDS database does not enforce TLS connections (rds.force_ssl / require_secure_transport)",
      severity=CAT_II, nist=["SC-8", "SC-8(1)"], area="encryption_transit",
      applies_to=["AWS::RDS::DBInstance", "AWS::RDS::DBCluster"], tags={"datastore"},
      description="The database has no parameter group in the template that sets rds.force_ssl=1 (PostgreSQL/SQL Server) or require_secure_transport=ON (MySQL/MariaDB). Clients may connect without TLS.",
      remediation="Create an AWS::RDS::DBParameterGroup/DBClusterParameterGroup that forces TLS and reference it from the database resource.")
def r_rds_force_ssl(ctx: TemplateCtx):
    pgs = dict(ctx.by_type("AWS::RDS::DBParameterGroup", "AWS::RDS::DBClusterParameterGroup"))
    for lid, res in ctx.by_type("AWS::RDS::DBInstance", "AWS::RDS::DBCluster"):
        p = props(res)
        if p.get("SourceDBInstanceIdentifier") or p.get("DBClusterIdentifier"):
            continue
        engine = (literal_str(p.get("Engine")) or "").lower()
        if engine and not re.search(r"postgres|mysql|mariadb|sqlserver|aurora", engine):
            continue
        ref = p.get("DBParameterGroupName") or p.get("DBClusterParameterGroupName")
        target = ref.get("Ref") if isinstance(ref, dict) else None
        if target in pgs and param_group_forces_ssl(pgs[target]):
            continue
        if target is None and isinstance(ref, str) and not ref.startswith("default."):
            continue  # external parameter group; cannot evaluate
        detail = "no DBParameterGroupName" if ref is None else f"parameter group {target or ref} does not force TLS"
        yield Hit(lid, rtype_of(res), ctx.resource_line(lid, res), detail)


@rule(id="CSA-TLS-005", title="Edge endpoint allows TLS below 1.2 or clear-text viewers",
      severity=CAT_II, nist=["SC-8(1)", "SC-13"], area="encryption_transit",
      applies_to=["AWS::ApiGateway::DomainName", "AWS::ApiGatewayV2::DomainName", "AWS::CloudFront::Distribution",
                  "AWS::Elasticsearch::Domain", "AWS::OpenSearchService::Domain", "AWS::ElastiCache::ReplicationGroup"],
      overlaps={"CKV_AWS_34", "CKV_AWS_174", "CKV_AWS_83", "CKV_AWS_30", "CKV_AWS_6", "CKV_AWS_228", "W70"}, tags={"exposure"},
      description="API Gateway custom domain SecurityPolicy is not TLS_1_2, CloudFront allows HTTP viewers or a MinimumProtocolVersion below TLSv1.2, OpenSearch does not enforce HTTPS with a TLS 1.2 policy, or ElastiCache does not enable in-transit encryption.",
      remediation="Set SecurityPolicy: TLS_1_2; ViewerProtocolPolicy redirect-to-https/https-only with MinimumProtocolVersion TLSv1.2_2021; DomainEndpointOptions EnforceHTTPS with Policy-Min-TLS-1-2; TransitEncryptionEnabled: true.")
def r_edge_tls(ctx: TemplateCtx):
    for lid, res in ctx.by_type("AWS::ApiGateway::DomainName", "AWS::ApiGatewayV2::DomainName"):
        p = props(res)
        pol = p.get("SecurityPolicy")
        if rtype_of(res) == "AWS::ApiGatewayV2::DomainName":
            confs = [c for c in as_list(p.get("DomainNameConfigurations")) if isinstance(c, dict)]
            pol = confs[0].get("SecurityPolicy") if confs else None
        if pol != "TLS_1_2":
            yield Hit(lid, rtype_of(res), ctx.resource_line(lid, res), f"SecurityPolicy={pol}")
    for lid, res in ctx.by_type("AWS::CloudFront::Distribution"):
        cfg = props(res).get("DistributionConfig") or {}
        if not isinstance(cfg, dict):
            continue
        dcb = cfg.get("DefaultCacheBehavior") or {}
        vpp = dcb.get("ViewerProtocolPolicy") if isinstance(dcb, dict) else None
        clear_text_paths = [str(b.get("PathPattern")) for b in as_list(cfg.get("CacheBehaviors"))
                            if isinstance(b, dict) and b.get("ViewerProtocolPolicy") == "allow-all"]
        cert = cfg.get("ViewerCertificate") or {}
        mpv_raw = cert.get("MinimumProtocolVersion") if isinstance(cert, dict) else None
        mpv_values = [v for v in ctx.resolve(mpv_raw) if v is not UNKNOWN] if mpv_raw is not None else []
        weak = [str(v) for v in mpv_values if not str(v).startswith(("TLSv1.2", "TLSv1.3"))]
        default_cert = isinstance(cert, dict) and is_true(cert.get("CloudFrontDefaultCertificate"))
        custom_cert = isinstance(cert, dict) and ("AcmCertificateArn" in cert or "IamCertificateId" in cert)
        if vpp == "allow-all":
            yield Hit(lid, rtype_of(res), ctx.resource_line(lid, res), "DefaultCacheBehavior.ViewerProtocolPolicy=allow-all")
        elif clear_text_paths:
            yield Hit(lid, rtype_of(res), ctx.resource_line(lid, res),
                      "CacheBehaviors ViewerProtocolPolicy=allow-all for PathPattern " + ", ".join(clear_text_paths))
        elif default_cert or not cert:
            yield Hit(lid, rtype_of(res), ctx.resource_line(lid, res), "ViewerCertificate MinimumProtocolVersion=default (TLSv1 with the CloudFront default certificate)")
        elif custom_cert and mpv_raw is None:
            yield Hit(lid, rtype_of(res), ctx.resource_line(lid, res), "ViewerCertificate has a custom certificate but no MinimumProtocolVersion (service default applies)")
        elif weak:
            via = f" (parameter {ctx.param_for(mpv_raw)} default)" if ctx.param_for(mpv_raw) else ""
            yield Hit(lid, rtype_of(res), ctx.resource_line(lid, res), f"ViewerCertificate MinimumProtocolVersion={weak[0]}{via}")
    for lid, res in ctx.by_type("AWS::Elasticsearch::Domain", "AWS::OpenSearchService::Domain"):
        o = props(res).get("DomainEndpointOptions") or {}
        if not (isinstance(o, dict) and all_true(ctx, o.get("EnforceHTTPS")) and str(o.get("TLSSecurityPolicy", "")).startswith("Policy-Min-TLS-1-2")):
            yield Hit(lid, rtype_of(res), ctx.resource_line(lid, res), "DomainEndpointOptions does not enforce HTTPS with a TLS 1.2 policy")
    for lid, res in ctx.by_type("AWS::ElastiCache::ReplicationGroup"):
        if not all_true(ctx, props(res).get("TransitEncryptionEnabled")):
            yield Hit(lid, rtype_of(res), ctx.resource_line(lid, res), "TransitEncryptionEnabled is not true")


# ---- Public exposure ----------------------------------------------------------------------


def port_range(rule_: dict) -> tuple[int | None, int | None, str]:
    proto = str(rule_.get("IpProtocol", "")).lower()
    try:
        f = int(rule_.get("FromPort")) if rule_.get("FromPort") is not None else None
        t = int(rule_.get("ToPort")) if rule_.get("ToPort") is not None else None
    except (TypeError, ValueError):
        f, t = None, None
    return f, t, proto


def ports_covered(f: int | None, t: int | None, proto: str, wanted: set[int]) -> bool:
    if proto in {"-1", "all"}:
        return True
    if proto not in {"tcp", "udp", "6", "17"} or f is None or t is None:
        return False
    if f == -1 or t == -1:
        return True
    return any(f <= p <= t for p in wanted)


def is_all_ports(f: int | None, t: int | None, proto: str) -> bool:
    if proto in {"icmp", "icmpv6", "1", "58"}:
        return False
    return proto in {"-1", "all"} or (f is not None and t is not None and ((f <= 0 and t >= 65535) or f == -1))


def iter_ingress(ctx: TemplateCtx) -> Iterable[tuple[str, str, dict, int]]:
    for lid, res in ctx.by_type("AWS::EC2::SecurityGroup"):
        for r in as_list(props(res).get("SecurityGroupIngress")):
            if isinstance(r, dict):
                yield lid, rtype_of(res), r, node_line(r, ctx.prop_line(lid, res, "Properties", "SecurityGroupIngress"))
    for lid, res in ctx.by_type("AWS::EC2::SecurityGroupIngress"):
        yield lid, rtype_of(res), props(res), ctx.resource_line(lid, res)


def public_cidr_source(ctx: TemplateCtx, r: dict) -> tuple[str | None, str | None]:
    """(cidr, parameter name) when the ingress source is 0.0.0.0/0 or ::/0."""
    for key in ("CidrIp", "CidrIpv6"):
        v = r.get(key)
        if v is None:
            continue
        for lit in ctx.resolve(v):
            if isinstance(lit, str) and lit.strip() in PUBLIC_CIDRS:
                return lit.strip(), ctx.param_for(v)
    return None, None


def _ingress_hits(ctx: TemplateCtx, wanted: set[int] | None, all_ports_only: bool = False):
    for lid, rtype, r, line in iter_ingress(ctx):
        cidr, param = public_cidr_source(ctx, r)
        if not cidr:
            continue
        f, t, proto = port_range(r)
        if all_ports_only:
            if not is_all_ports(f, t, proto):
                continue
        elif is_all_ports(f, t, proto) or not ports_covered(f, t, proto, wanted or set()):
            continue
        yield lid, rtype, r, line, cidr, param, f, t, proto


def ingress_key(cidr: str, f: int | None, t: int | None, proto: str) -> str:
    """Content-based identity of an ingress rule, independent of its position among sibling rules."""
    return f"{cidr}|{f}-{t}/{proto}"


def _ipv6_cis(cidr: str) -> str:
    return CIS["sg_admin_ipv6"] if cidr == "::/0" else CIS["sg_admin_ipv4"]


@rule(id="CSA-NET-001", title="Security group allows unrestricted ingress to administrative ports",
      severity=CAT_I, nist=["SC-7", "SC-7(5)", "AC-17"], area="boundary", cis=CIS["sg_admin_ipv4"],
      applies_to=["AWS::EC2::SecurityGroup", "AWS::EC2::SecurityGroupIngress"],
      overlaps={"CKV_AWS_24", "CKV_AWS_25", "W2", "W9", "W40"}, tags={"exposure"},
      description="SSH (22), RDP (3389) or WinRM (5985/5986) is reachable from 0.0.0.0/0 or ::/0. Any host on the Internet can attempt to authenticate to the instance.",
      remediation="Replace the open CIDR with a parameterized Government CIDR (AllowedPattern that rejects /0), or remove the rule and use Systems Manager Session Manager / a bastion behind the CAP.")
def r_sg_admin(ctx: TemplateCtx):
    for lid, rtype, _r, line, cidr, param, f, t, proto in _ingress_hits(ctx, ADMIN_PORTS):
        if param:
            yield Hit(lid, rtype, line, f"ports {f}-{t}/{proto} from parameter {param} whose Default is {cidr}", severity=CAT_II,
                      title="Security group admin-port ingress defaults to an unrestricted CIDR parameter", key=ingress_key(cidr, f, t, proto))
        else:
            yield Hit(lid, rtype, line, f"ports {f}-{t}/{proto} from {cidr}", key=ingress_key(cidr, f, t, proto))


@rule(id="CSA-NET-002", title="Security group allows unrestricted ingress to database or cache ports",
      severity=CAT_I, nist=["SC-7", "SC-7(5)", "AC-3"], area="boundary",
      applies_to=["AWS::EC2::SecurityGroup", "AWS::EC2::SecurityGroupIngress"], overlaps={"W2", "W9"}, tags={"exposure", "datastore"},
      description="A database, cache or search port is reachable from 0.0.0.0/0 or ::/0. Data-tier services must only accept traffic from application-tier security groups.",
      remediation="Replace the CIDR source with SourceSecurityGroupId of the application tier, or a parameterized Government CIDR.")
def r_sg_db(ctx: TemplateCtx):
    for lid, rtype, _r, line, cidr, param, f, t, proto in _ingress_hits(ctx, DATABASE_PORTS):
        if ports_covered(f, t, proto, ADMIN_PORTS):
            continue  # reported by CSA-NET-001
        detail = f"ports {f}-{t}/{proto} from {cidr}" + (f" (parameter {param} default)" if param else "")
        yield Hit(lid, rtype, line, detail, severity=CAT_II if param else CAT_I, key=ingress_key(cidr, f, t, proto))


@rule(id="CSA-NET-003", title="Security group allows unrestricted ingress on all ports and protocols",
      severity=CAT_I, nist=["SC-7", "SC-7(5)", "CM-7"], area="boundary",
      applies_to=["AWS::EC2::SecurityGroup", "AWS::EC2::SecurityGroupIngress"], overlaps={"W2", "W9", "W40", "W42"}, tags={"exposure"},
      description="An ingress rule with IpProtocol -1 or the full port range is open to 0.0.0.0/0 or ::/0.",
      remediation="Restrict the rule to the required protocol and ports and to a known source security group or Government CIDR.")
def r_sg_all(ctx: TemplateCtx):
    for lid, rtype, _r, line, cidr, param, f, t, proto in _ingress_hits(ctx, None, all_ports_only=True):
        yield Hit(lid, rtype, line, f"IpProtocol={proto} ports {f}-{t} from {cidr}" + (f" (parameter {param} default)" if param else ""),
                  severity=CAT_II if param else CAT_I, key=ingress_key(cidr, f, t, proto))


@rule(id="CSA-NET-004", title="Security group allows unrestricted ingress on an application port",
      severity=CAT_III, nist=["SC-7", "SC-7(5)"], area="boundary",
      applies_to=["AWS::EC2::SecurityGroup", "AWS::EC2::SecurityGroupIngress"], overlaps={"W2", "W9", "CKV_AWS_260"}, tags={"exposure"},
      description="A non-administrative port is open to 0.0.0.0/0 or ::/0. This is expected for a public web tier behind the Cloud Access Point, but each rule needs Government confirmation of intent.",
      remediation="Confirm the resource is an Internet-facing tier. Front it with a WAF and the CAP/VDSS path; otherwise restrict the source to the load balancer security group.")
def r_sg_app(ctx: TemplateCtx):
    for lid, rtype, r, line in iter_ingress(ctx):
        cidr, param = public_cidr_source(ctx, r)
        if not cidr:
            continue
        f, t, proto = port_range(r)
        if is_all_ports(f, t, proto) or ports_covered(f, t, proto, ADMIN_PORTS) or ports_covered(f, t, proto, DATABASE_PORTS):
            continue
        web = ports_covered(f, t, proto, WEB_PORTS) and f == t
        yield Hit(lid, rtype, line, f"ports {f}-{t}/{proto} from {cidr}" + (f" (parameter {param} default)" if param else ""),
                  severity=CAT_III if web else CAT_II, disposition="Risk acceptance recommended" if web else None, key=ingress_key(cidr, f, t, proto))


@rule(id="CSA-NET-005", title="Database is publicly accessible",
      severity=CAT_I, nist=["SC-7", "AC-3"], area="boundary", cis=CIS["rds_public"],
      applies_to=["AWS::RDS::DBInstance", "AWS::Redshift::Cluster", "AWS::DocDB::DBInstance"], overlaps={"CKV_AWS_17", "CKV_AWS_87", "F22"},
      tags={"exposure", "datastore"},
      description="PubliclyAccessible is true, so the database receives a public IP address and DNS name reachable from the Internet.",
      remediation="Set PubliclyAccessible: false and place the database in private subnets reachable only from the application tier.")
def r_db_public(ctx: TemplateCtx):
    for lid, res in ctx.by_type("AWS::RDS::DBInstance", "AWS::Redshift::Cluster", "AWS::DocDB::DBInstance"):
        v = props(res).get("PubliclyAccessible")
        if any_true(ctx, v):
            yield Hit(lid, rtype_of(res), ctx.prop_line(lid, res, "Properties", "PubliclyAccessible"), f"PubliclyAccessible resolves to {ctx.resolve(v)}")


PUBLIC_ACLS = {"PublicRead", "PublicReadWrite", "AuthenticatedRead"}


@rule(id="CSA-NET-006", title="S3 bucket grants public or anonymous access",
      severity=CAT_I, nist=["AC-3", "SC-7", "AC-6"], area="boundary",
      applies_to=[S3_BUCKET, "AWS::S3::BucketPolicy"], overlaps={"CKV_AWS_20", "CKV_AWS_57", "CKV_AWS_70", "F14", "F15", "F16"},
      tags={"exposure", "datastore"},
      description="The bucket ACL is PublicRead/PublicReadWrite/AuthenticatedRead, or a bucket policy allows Principal * without a restricting condition. Any Internet user can read (or write) objects.",
      remediation="Remove the public ACL, enable PublicAccessBlockConfiguration, and serve public content through CloudFront with origin access control instead of a public bucket.")
def r_s3_public(ctx: TemplateCtx):
    for lid, res in ctx.by_type(S3_BUCKET):
        acl = props(res).get("AccessControl")
        if isinstance(acl, str) and acl in PUBLIC_ACLS:
            yield Hit(lid, S3_BUCKET, ctx.prop_line(lid, res, "Properties", "AccessControl"), f"AccessControl={acl}")
    for lid, res in ctx.by_type("AWS::S3::BucketPolicy"):
        for s in iter_statements(props(res).get("PolicyDocument")):
            if s.get("Effect") == "Allow" and principal_is_public(s.get("Principal")) and not statement_has_condition(s):
                yield Hit(lid, rtype_of(res), node_line(s, ctx.resource_line(lid, res)), f"Allow statement with Principal * for {stmt_actions(s)}")
                break


@rule(id="CSA-NET-007", title="S3 bucket does not block public access",
      severity=CAT_II, nist=["AC-3", "SC-7", "CM-6"], area="boundary", cis=CIS["s3_block_public"],
      applies_to=[S3_BUCKET], overlaps={"CKV_AWS_53", "CKV_AWS_54", "CKV_AWS_55", "CKV_AWS_56"}, tags={"exposure", "datastore"},
      description="PublicAccessBlockConfiguration is absent or does not set all four settings to true. A later ACL or policy change can expose the bucket.",
      remediation="Add PublicAccessBlockConfiguration with BlockPublicAcls, BlockPublicPolicy, IgnorePublicAcls and RestrictPublicBuckets all true.")
def r_s3_pab(ctx: TemplateCtx):
    keys = ("BlockPublicAcls", "BlockPublicPolicy", "IgnorePublicAcls", "RestrictPublicBuckets")
    for lid, res in ctx.by_type(S3_BUCKET):
        pab = props(res).get("PublicAccessBlockConfiguration")
        if not isinstance(pab, dict):
            yield Hit(lid, S3_BUCKET, ctx.resource_line(lid, res), "PublicAccessBlockConfiguration is absent")
        else:
            missing = [k for k in keys if not all_true(ctx, pab.get(k))]
            if missing:
                yield Hit(lid, S3_BUCKET, ctx.prop_line(lid, res, "Properties", "PublicAccessBlockConfiguration"), f"not true: {', '.join(missing)}")


@rule(id="CSA-NET-008", title="EKS cluster API endpoint is reachable from the Internet",
      severity=CAT_II, nist=["SC-7", "AC-17"], area="boundary",
      applies_to=["AWS::EKS::Cluster"], overlaps={"CKV_AWS_38", "CKV_AWS_39"}, tags={"exposure"},
      description="ResourcesVpcConfig.EndpointPublicAccess is not false and PublicAccessCidrs is absent or 0.0.0.0/0. The Kubernetes API is exposed to any Internet source.",
      remediation="Set EndpointPrivateAccess: true and EndpointPublicAccess: false, or restrict PublicAccessCidrs to Government CIDRs.")
def r_eks_public(ctx: TemplateCtx):
    for lid, res in ctx.by_type("AWS::EKS::Cluster"):
        vpc = props(res).get("ResourcesVpcConfig") or {}
        if not isinstance(vpc, dict):
            continue
        pub = vpc.get("EndpointPublicAccess")
        if pub is not None and not any_true(ctx, pub) and UNKNOWN not in ctx.resolve(pub):
            continue
        cidrs = [c for c in as_list(vpc.get("PublicAccessCidrs")) if isinstance(c, str)]
        if cidrs and not any(c in PUBLIC_CIDRS for c in cidrs):
            continue
        yield Hit(lid, rtype_of(res), ctx.prop_line(lid, res, "Properties", "ResourcesVpcConfig"),
                  f"EndpointPublicAccess={pub if pub is not None else 'default true'}, PublicAccessCidrs={cidrs or 'default 0.0.0.0/0'}")


def public_subnets(ctx: TemplateCtx) -> set[str]:
    pub: set[str] = set()
    igws = {lid for lid, _ in ctx.by_type("AWS::EC2::InternetGateway")}
    public_rts: set[str] = set()
    for lid, res in ctx.by_type("AWS::EC2::Route"):
        p = props(res)
        gw = p.get("GatewayId")
        dest = p.get("DestinationCidrBlock") or p.get("DestinationIpv6CidrBlock")
        if isinstance(dest, str) and dest in PUBLIC_CIDRS and isinstance(gw, dict) and isinstance(gw.get("Ref"), str) and gw["Ref"] in igws:
            rt = p.get("RouteTableId")
            if isinstance(rt, dict) and rt.get("Ref"):
                public_rts.add(rt["Ref"])
    for lid, res in ctx.by_type("AWS::EC2::SubnetRouteTableAssociation"):
        p = props(res)
        rt, sn = p.get("RouteTableId"), p.get("SubnetId")
        if isinstance(rt, dict) and rt.get("Ref") in public_rts and isinstance(sn, dict) and sn.get("Ref"):
            pub.add(sn["Ref"])
    for lid, res in ctx.by_type("AWS::EC2::Subnet"):
        if any_true(ctx, props(res).get("MapPublicIpOnLaunch")):
            pub.add(lid)
    return pub


@rule(id="CSA-NET-009", title="Data-tier subnet group uses public subnets",
      severity=CAT_II, nist=["SC-7", "AC-4"], area="boundary",
      applies_to=["AWS::RDS::DBSubnetGroup", "AWS::ElastiCache::SubnetGroup", "AWS::Redshift::ClusterSubnetGroup", "AWS::DocDB::DBSubnetGroup"],
      tags={"exposure", "datastore"},
      description="The subnet group references a subnet that has a default route to an Internet gateway or MapPublicIpOnLaunch: true. Databases should live in private subnets.",
      remediation="Create dedicated private subnets (no 0.0.0.0/0 route to an Internet gateway) for the data tier and reference them from the subnet group.")
def r_db_public_subnet(ctx: TemplateCtx):
    pub = public_subnets(ctx)
    if not pub:
        return
    for lid, res in ctx.by_type("AWS::RDS::DBSubnetGroup", "AWS::ElastiCache::SubnetGroup", "AWS::Redshift::ClusterSubnetGroup", "AWS::DocDB::DBSubnetGroup"):
        refs = [s.get("Ref") for s in as_list(props(res).get("SubnetIds")) if isinstance(s, dict) and isinstance(s.get("Ref"), str)]
        bad = [r for r in refs if r in pub]
        if bad:
            yield Hit(lid, rtype_of(res), ctx.prop_line(lid, res, "Properties", "SubnetIds"), f"public subnets: {', '.join(bad)}")


@rule(id="CSA-NET-010", title="API method or function URL has no authorization",
      severity=CAT_II, nist=["AC-3", "IA-2", "SC-7"], area="boundary",
      applies_to=["AWS::ApiGateway::Method", "AWS::Lambda::Url", "AWS::ApiGatewayV2::Route"], overlaps={"CKV_AWS_59", "CKV_AWS_258", "CKV_AWS_309"},
      tags={"exposure"},
      description="AuthorizationType is NONE (and ApiKeyRequired is not true) or the Lambda function URL uses AuthType NONE. The endpoint is callable by anyone who can reach it.",
      remediation="Use AWS_IAM, a Cognito user pool or a Lambda authorizer integrated with the Government identity provider; require an API key and usage plan at minimum.")
def r_api_noauth(ctx: TemplateCtx):
    for lid, res in ctx.by_type("AWS::ApiGateway::Method"):
        p = props(res)
        if p.get("HttpMethod") == "OPTIONS":
            continue
        if str(p.get("AuthorizationType", "NONE")).upper() == "NONE" and not any_true(ctx, p.get("ApiKeyRequired")):
            yield Hit(lid, rtype_of(res), ctx.prop_line(lid, res, "Properties", "AuthorizationType"), f"AuthorizationType={p.get('AuthorizationType', 'NONE')}")
    for lid, res in ctx.by_type("AWS::ApiGatewayV2::Route"):
        p = props(res)
        if str(p.get("AuthorizationType", "NONE")).upper() == "NONE" and not any_true(ctx, p.get("ApiKeyRequired")):
            yield Hit(lid, rtype_of(res), ctx.resource_line(lid, res), f"AuthorizationType={p.get('AuthorizationType', 'NONE')}")
    for lid, res in ctx.by_type("AWS::Lambda::Url"):
        if str(props(res).get("AuthType", "")).upper() == "NONE":
            yield Hit(lid, rtype_of(res), ctx.prop_line(lid, res, "Properties", "AuthType"), "AuthType=NONE")


# ---- Logging and monitoring ----------------------------------------------------------------


@rule(id="CSA-LOG-001", title="S3 bucket has no server access logging",
      severity=CAT_III, nist=["AU-2", "AU-12"], area="logging",
      applies_to=[S3_BUCKET], overlaps={"CKV_AWS_18", "W35"},
      description="LoggingConfiguration is absent. Object-level access to the bucket cannot be reconstructed for audit or incident response.",
      remediation="Add LoggingConfiguration pointing at a dedicated, encrypted log bucket, or enable CloudTrail data events for the bucket.")
def r_s3_logging(ctx: TemplateCtx):
    for lid, res in ctx.by_type(S3_BUCKET):
        if "LoggingConfiguration" not in props(res):
            yield Hit(lid, S3_BUCKET, ctx.resource_line(lid, res), "Properties.LoggingConfiguration is absent")


@rule(id="CSA-LOG-002", title="VPC has no flow log",
      severity=CAT_II, nist=["AU-2", "AU-12", "SI-4"], area="logging", cis=CIS["vpc_flow_logs"],
      applies_to=["AWS::EC2::VPC"], overlaps={"W60"},
      description="The template creates a VPC but no AWS::EC2::FlowLog references it. Network traffic metadata is not captured for monitoring or forensics.",
      remediation="Add an AWS::EC2::FlowLog (ResourceType VPC, TrafficType ALL) delivering to an encrypted CloudWatch log group or S3 bucket.")
def r_vpc_flow_logs(ctx: TemplateCtx):
    logged = set()
    for lid, res in ctx.by_type("AWS::EC2::FlowLog"):
        rid = props(res).get("ResourceId")
        if isinstance(rid, dict) and rid.get("Ref"):
            logged.add(rid["Ref"])
        elif isinstance(rid, dict) and "Fn::GetAtt" in rid:
            logged.add(as_list(rid["Fn::GetAtt"])[0])
    for lid, res in ctx.by_type("AWS::EC2::VPC"):
        if lid not in logged:
            yield Hit(lid, "AWS::EC2::VPC", ctx.resource_line(lid, res), "no AWS::EC2::FlowLog references this VPC")


@rule(id="CSA-LOG-003", title="Load balancer access logging is disabled",
      severity=CAT_III, nist=["AU-2", "AU-12"], area="logging",
      applies_to=["AWS::ElasticLoadBalancingV2::LoadBalancer", "AWS::ElasticLoadBalancing::LoadBalancer"], overlaps={"CKV_AWS_91", "CKV_AWS_92", "W26", "W52"},
      description="access_logs.s3.enabled is not true (ALB/NLB) or AccessLoggingPolicy is not enabled (Classic ELB). Request-level records are not retained.",
      remediation="Enable access logs to an encrypted S3 bucket with a bucket policy for the regional ELB service account.")
def r_lb_logs(ctx: TemplateCtx):
    for lid, res in ctx.by_type("AWS::ElasticLoadBalancingV2::LoadBalancer"):
        attrs = {a.get("Key"): a.get("Value") for a in as_list(props(res).get("LoadBalancerAttributes")) if isinstance(a, dict)}
        if not any_true(ctx, attrs.get("access_logs.s3.enabled")):
            yield Hit(lid, rtype_of(res), ctx.resource_line(lid, res), "LoadBalancerAttributes access_logs.s3.enabled is not true")
    for lid, res in ctx.by_type("AWS::ElasticLoadBalancing::LoadBalancer"):
        pol = props(res).get("AccessLoggingPolicy")
        if not (isinstance(pol, dict) and any_true(ctx, pol.get("Enabled"))):
            yield Hit(lid, rtype_of(res), ctx.resource_line(lid, res), "AccessLoggingPolicy is absent or disabled")


@rule(id="CSA-LOG-004", title="API Gateway stage has no execution or access logging",
      severity=CAT_III, nist=["AU-2", "AU-12"], area="logging",
      applies_to=["AWS::ApiGateway::Stage", "AWS::ApiGatewayV2::Stage", "AWS::Serverless::Api"], overlaps={"CKV_AWS_76", "CKV_AWS_95"},
      description="The stage has no MethodSettings LoggingLevel (INFO/ERROR) and no AccessLogSetting. API calls are not recorded.",
      remediation="Add AccessLogSetting to a CloudWatch log group and MethodSettings with LoggingLevel: INFO and DataTraceEnabled: false.")
def r_apigw_logs(ctx: TemplateCtx):
    for lid, res in ctx.by_type("AWS::ApiGateway::Stage", "AWS::Serverless::Api"):
        p = props(res)
        has_exec = any(isinstance(m, dict) and str(m.get("LoggingLevel", "OFF")).upper() in {"INFO", "ERROR"} for m in as_list(p.get("MethodSettings")))
        if not has_exec and not p.get("AccessLogSetting"):
            yield Hit(lid, rtype_of(res), ctx.resource_line(lid, res), "no MethodSettings.LoggingLevel and no AccessLogSetting")
    for lid, res in ctx.by_type("AWS::ApiGatewayV2::Stage"):
        if not props(res).get("AccessLogSettings"):
            yield Hit(lid, rtype_of(res), ctx.resource_line(lid, res), "AccessLogSettings is absent")


@rule(id="CSA-LOG-005", title="RDS database does not export logs to CloudWatch Logs",
      severity=CAT_III, nist=["AU-2", "AU-12"], area="logging",
      applies_to=["AWS::RDS::DBInstance", "AWS::RDS::DBCluster"], overlaps={"CKV_AWS_129"},
      description="EnableCloudwatchLogsExports is absent. Database error, audit and slow-query logs stay on the instance and are not centrally retained.",
      remediation="Set EnableCloudwatchLogsExports for the engine (for example [postgresql, upgrade] or [audit, error, general, slowquery]).")
def r_rds_logs(ctx: TemplateCtx):
    for lid, res in ctx.by_type("AWS::RDS::DBInstance", "AWS::RDS::DBCluster"):
        p = props(res)
        if p.get("SourceDBInstanceIdentifier") or p.get("DBClusterIdentifier"):
            continue
        if not p.get("EnableCloudwatchLogsExports"):
            yield Hit(lid, rtype_of(res), ctx.resource_line(lid, res), "EnableCloudwatchLogsExports is absent")


@rule(id="CSA-LOG-006", title="EKS control-plane logging is incomplete",
      severity=CAT_II, nist=["AU-2", "AU-12"], area="logging",
      applies_to=["AWS::EKS::Cluster"], overlaps={"CKV_AWS_37"},
      description="Logging.ClusterLogging does not enable the api, audit and authenticator log types. Kubernetes API and authentication events are not retained.",
      remediation="Set Logging.ClusterLogging.EnabledTypes to api, audit, authenticator, controllerManager and scheduler.")
def r_eks_logs(ctx: TemplateCtx):
    for lid, res in ctx.by_type("AWS::EKS::Cluster"):
        logging_ = props(res).get("Logging") or {}
        types: set[str] = set()
        cl = logging_.get("ClusterLogging") if isinstance(logging_, dict) else None
        for config in as_list(cl):  # CloudFormation models ClusterLogging as a list of {EnabledTypes: [...]}
            if not isinstance(config, dict):
                continue
            for e in as_list(config.get("EnabledTypes")):
                if isinstance(e, dict) and isinstance(e.get("Type"), str):
                    types.add(e["Type"])
        missing = {"api", "audit", "authenticator"} - types
        if missing:
            yield Hit(lid, rtype_of(res), ctx.resource_line(lid, res), f"missing control-plane log types: {', '.join(sorted(missing))}")


@rule(id="CSA-LOG-007", title="CloudTrail trail is not multi-Region or lacks log file validation",
      severity=CAT_III, nist=["AU-2", "AU-9", "AU-12"], area="logging", cis=CIS["cloudtrail_multiregion"],
      applies_to=["AWS::CloudTrail::Trail"], overlaps={"CKV_AWS_36", "CKV_AWS_67"},
      description="IsMultiRegionTrail is not true (CIS 3.1) or EnableLogFileValidation is not true (CIS 3.2). Events in other Regions are missed or log tampering cannot be detected.",
      remediation="Set IsMultiRegionTrail: true, IncludeGlobalServiceEvents: true and EnableLogFileValidation: true.")
def r_cloudtrail_cfg(ctx: TemplateCtx):
    for lid, res in ctx.by_type("AWS::CloudTrail::Trail"):
        p = props(res)
        issues = []
        if not all_true(ctx, p.get("IsMultiRegionTrail")):
            issues.append("IsMultiRegionTrail is not true")
        if not all_true(ctx, p.get("EnableLogFileValidation")):
            issues.append("EnableLogFileValidation is not true")
        if issues:
            yield Hit(lid, rtype_of(res), ctx.resource_line(lid, res), "; ".join(issues))


@rule(id="CSA-LOG-008", title="CloudWatch log group has no retention period",
      severity=CAT_III, nist=["AU-11", "SI-12"], area="logging",
      applies_to=["AWS::Logs::LogGroup"], overlaps={"CKV_AWS_66", "W86"},
      description="RetentionInDays is absent, so logs are kept forever without a documented retention decision.",
      remediation="Set RetentionInDays to the Government-approved retention period (for example 365 or 400 days) or archive to S3 with lifecycle rules.")
def r_loggroup_retention(ctx: TemplateCtx):
    for lid, res in ctx.by_type("AWS::Logs::LogGroup"):
        if "RetentionInDays" not in props(res):
            yield Hit(lid, rtype_of(res), ctx.resource_line(lid, res), "RetentionInDays is absent")


@rule(id="CSA-MON-001", title="Lambda function has no active tracing",
      severity=CAT_III, nist=["SI-4", "AU-12"], area="monitoring",
      applies_to=["AWS::Lambda::Function", "AWS::Serverless::Function"], overlaps={"CKV_AWS_50"},
      description="TracingConfig.Mode is not Active. Invocation traces are not available for anomaly detection or incident reconstruction.",
      remediation="Set TracingConfig: {Mode: Active} (Tracing: Active for SAM) and grant xray:PutTraceSegments.")
def r_lambda_tracing(ctx: TemplateCtx):
    for lid, res in ctx.by_type("AWS::Lambda::Function", "AWS::Serverless::Function"):
        p = props(res)
        tc = p.get("TracingConfig")
        mode = tc.get("Mode") if isinstance(tc, dict) else p.get("Tracing")
        if mode != "Active":
            yield Hit(lid, rtype_of(res), ctx.resource_line(lid, res), "TracingConfig.Mode is not Active")


@rule(id="CSA-MON-002", title="Lambda function is not attached to a VPC although the stack contains VPC data resources",
      severity=CAT_III, nist=["SC-7", "AC-4"], area="boundary",
      applies_to=["AWS::Lambda::Function", "AWS::Serverless::Function"], overlaps={"CKV_AWS_117", "W89"},
      description="The function has no VpcConfig while the same template defines a VPC, database, cache or file system. Traffic from the function does not pass through the VPC security controls.",
      remediation="Add VpcConfig with private subnets and a dedicated security group; use VPC endpoints for AWS service access.")
def r_lambda_vpc(ctx: TemplateCtx):
    if not ctx.has_type("AWS::EC2::VPC", "AWS::RDS::DBInstance", "AWS::RDS::DBCluster", "AWS::EFS::FileSystem",
                        "AWS::ElastiCache::ReplicationGroup", "AWS::ElastiCache::CacheCluster"):
        return
    for lid, res in ctx.by_type("AWS::Lambda::Function", "AWS::Serverless::Function"):
        if "VpcConfig" not in props(res):
            yield Hit(lid, rtype_of(res), ctx.resource_line(lid, res), "VpcConfig is absent")


# ---- IAM least privilege --------------------------------------------------------------------

READ_ONLY_PREFIXES = ("get", "list", "describe", "head", "lookup", "search", "query", "scan", "batchget", "select", "view", "check", "detect")
IAM_POLICY_DOCUMENT_TYPES = {"AWS::IAM::Policy", "AWS::IAM::ManagedPolicy", "AWS::IAM::RolePolicy", "AWS::IAM::UserPolicy", "AWS::IAM::GroupPolicy"}
IAM_PRINCIPAL_TYPES = sorted(IAM_POLICY_DOCUMENT_TYPES) + ["AWS::IAM::Role", "AWS::IAM::User", "AWS::IAM::Group"]


def iter_policy_documents(ctx: TemplateCtx) -> Iterable[tuple[str, str, dict, int, str, str]]:
    """Yield (logical_id, resource_type, statement, line, context, statement key) for identity-based policies.

    The statement key is ``<context>.Statement[<Sid>]`` when the statement has a Sid, otherwise a digest of the
    statement's Effect/Action/Resource, so the key does not move when sibling statements are added or removed."""
    for lid, res in ctx.resources.items():
        rtype = rtype_of(res)
        p = props(res)
        docs: list[tuple[Any, str]] = []
        if rtype in IAM_POLICY_DOCUMENT_TYPES:
            docs.append((p.get("PolicyDocument"), "PolicyDocument"))
        if rtype in {"AWS::IAM::Role", "AWS::IAM::User", "AWS::IAM::Group"}:
            for i, pol in enumerate(as_list(p.get("Policies"))):
                if isinstance(pol, dict):
                    name = pol.get("PolicyName")
                    docs.append((pol.get("PolicyDocument"), f"Policies[{name if isinstance(name, str) and name else i}]"))
        for doc, where in docs:
            for s in iter_statements(doc):
                sid = s.get("Sid")
                if not (isinstance(sid, str) and sid):
                    body = json.dumps(_canon({k: s.get(k) for k in ("Effect", "Action", "NotAction", "Resource", "NotResource")}), sort_keys=True, default=str)
                    sid = hashlib.sha1(body.encode()).hexdigest()[:8]
                key = f"{where}.Statement[{sid}]"
                yield lid, rtype, s, node_line(s, ctx.resource_line(lid, res)), where, key


def is_read_only(actions: list[str]) -> bool:
    for a in actions:
        op = a.split(":", 1)[1].lower() if ":" in a else a.lower()
        if not op.startswith(READ_ONLY_PREFIXES):
            return False
    return True


@rule(id="CSA-IAM-001", title="IAM policy grants full administrative privileges (Action * on Resource *)",
      severity=CAT_I, nist=["AC-6", "AC-6(1)", "AC-3"], area="iam",
      applies_to=IAM_PRINCIPAL_TYPES,
      overlaps={"CKV_AWS_1", "CKV_AWS_107", "CKV_AWS_108", "CKV_AWS_109", "CKV_AWS_110", "CKV_AWS_111", "CKV_AWS_286", "CKV_AWS_287", "CKV_AWS_288",
                "CKV_AWS_289", "CKV_AWS_290", "F2", "F3", "F4", "F5", "F38", "F39", "F40", "F41"},
      description="An Allow statement uses Action \"*\" with Resource \"*\". The principal can perform any action on any resource in the account, including privilege escalation and log tampering.",
      remediation="Enumerate the required actions and scope Resource to specific ARNs; attach a permissions boundary to the role.")
def r_iam_admin(ctx: TemplateCtx):
    for lid, rtype, s, line, where, key in iter_policy_documents(ctx):
        if s.get("Effect") != "Allow":
            continue
        if "*" in stmt_actions(s) and any(r == "*" for r in stmt_resources(s)):
            yield Hit(lid, rtype, line, f"{where}: Action * on Resource *", key=key)


@rule(id="CSA-IAM-002", title="IAM policy uses service-level wildcard actions",
      severity=CAT_II, nist=["AC-6", "AC-6(1)"], area="iam",
      applies_to=IAM_PRINCIPAL_TYPES, overlaps={"CKV_AWS_107", "CKV_AWS_108", "CKV_AWS_109", "CKV_AWS_110", "CKV_AWS_111", "F2", "F3", "F4", "F5", "W11", "W12", "W13"},
      description="An Allow statement grants service:* (for example s3:* or ec2:*). The principal receives every current and future action for the service, far beyond what the workload needs.",
      remediation="Replace service:* with the specific actions the workload calls; use IAM Access Analyzer policy generation from CloudTrail to derive the list.")
def r_iam_service_wildcard(ctx: TemplateCtx):
    for lid, rtype, s, line, where, key in iter_policy_documents(ctx):
        if s.get("Effect") != "Allow":
            continue
        actions = stmt_actions(s)
        if "*" in actions and any(r == "*" for r in stmt_resources(s)):
            continue  # CSA-IAM-001
        wild = [a for a in actions if a.endswith(":*") or a == "*"]
        if wild:
            yield Hit(lid, rtype, line, f"{where}: {', '.join(wild[:4])}", key=key)


@rule(id="CSA-IAM-003", title="IAM policy allows write actions on Resource *",
      severity=CAT_II, nist=["AC-6", "AC-3"], area="iam",
      applies_to=IAM_PRINCIPAL_TYPES, overlaps={"CKV_AWS_356", "CKV_AWS_355", "W11", "W12", "W13", "F4", "F5", "F39"},
      description="An Allow statement with non-read-only actions applies to Resource \"*\". The principal can modify or delete any resource of that type in the account.",
      remediation="Scope Resource to the ARNs the workload owns (use !Sub with the stack's resource names) and add Condition keys where the service requires *.")
def r_iam_resource_wildcard(ctx: TemplateCtx):
    for lid, rtype, s, line, where, key in iter_policy_documents(ctx):
        if s.get("Effect") != "Allow":
            continue
        actions = stmt_actions(s)
        if not actions or "*" in actions or any(a.endswith(":*") for a in actions):
            continue  # reported by IAM-001/002
        if not any(r == "*" for r in stmt_resources(s)) or is_read_only(actions) or statement_has_condition(s):
            continue
        # Services / actions that only support Resource *
        if all(a.split(":")[0].lower() in {"cloudwatch", "xray", "ec2messages", "ssmmessages", "sts"}
               or a.lower() in {"logs:createloggroup", "logs:describeloggroups", "logs:describelogstreams"} for a in actions):
            continue
        yield Hit(lid, rtype, line, f"{where}: {', '.join(actions[:4])}{'...' if len(actions) > 4 else ''} on Resource *", key=key)


@rule(id="CSA-IAM-004", title="IAM policy allows iam:PassRole on Resource *",
      severity=CAT_II, nist=["AC-6", "AC-6(1)", "AC-3"], area="iam",
      applies_to=IAM_PRINCIPAL_TYPES, overlaps={"CKV_AWS_110"},
      description="iam:PassRole (or iam:*) is allowed on any role. The principal can pass a more privileged role to a service and escalate privileges.",
      remediation="Restrict iam:PassRole to the specific role ARNs and add Condition iam:PassedToService.")
def r_iam_passrole(ctx: TemplateCtx):
    for lid, rtype, s, line, where, key in iter_policy_documents(ctx):
        if s.get("Effect") != "Allow":
            continue
        actions = [a.lower() for a in stmt_actions(s)]
        if "*" in actions:
            continue
        if any(a in {"iam:passrole", "iam:*", "iam:pass*"} for a in actions) and any(r == "*" for r in stmt_resources(s)) and not statement_has_condition(s):
            yield Hit(lid, rtype, line, f"{where}: iam:PassRole on Resource *", key=key)


@rule(id="CSA-IAM-005", title="IAM user or group carries inline policies or long-lived credentials",
      severity=CAT_II, nist=["AC-2", "AC-6", "IA-5"], area="iam",
      applies_to=["AWS::IAM::User", "AWS::IAM::Group", "AWS::IAM::AccessKey"], overlaps={"CKV_AWS_40", "CKV_AWS_273", "CKV_AWS_274", "F10", "F11", "F12"},
      description="The template creates IAM users with inline policies, login profiles or access keys. Long-lived credentials bypass the Government identity provider and MFA, and inline policies are not reusable or reviewable.",
      remediation="Use federated roles through the Government identity provider; if a user is unavoidable, attach managed policies through groups and do not create access keys in templates.")
def r_iam_users(ctx: TemplateCtx):
    for lid, res in ctx.by_type("AWS::IAM::User"):
        p = props(res)
        issues = []
        if p.get("Policies"):
            issues.append("inline Policies")
        if p.get("LoginProfile"):
            issues.append("LoginProfile (console password)")
        if issues:
            yield Hit(lid, rtype_of(res), ctx.resource_line(lid, res), "; ".join(issues))
    for lid, res in ctx.by_type("AWS::IAM::AccessKey"):
        yield Hit(lid, rtype_of(res), ctx.resource_line(lid, res), "AWS::IAM::AccessKey creates a long-lived access key")
    for lid, res in ctx.by_type("AWS::IAM::Group"):
        if props(res).get("Policies"):
            yield Hit(lid, rtype_of(res), ctx.resource_line(lid, res), "inline Policies on group", severity=CAT_III)


@rule(id="CSA-IAM-006", title="IAM role uses inline policies",
      severity=CAT_III, nist=["AC-6", "CM-6"], area="iam",
      applies_to=["AWS::IAM::Role"],
      description="The role embeds policy documents in Policies. Inline policies cannot be versioned, reviewed or reused independently of the role.",
      remediation="Move the statements to an AWS::IAM::ManagedPolicy (customer managed) and reference it from ManagedPolicyArns.")
def r_iam_inline(ctx: TemplateCtx):
    for lid, res in ctx.by_type("AWS::IAM::Role"):
        if props(res).get("Policies"):
            yield Hit(lid, rtype_of(res), ctx.prop_line(lid, res, "Properties", "Policies"), "Properties.Policies is present")


@rule(id="CSA-IAM-007", title="IAM role has no permissions boundary",
      severity=CAT_III, nist=["AC-6", "AC-6(1)", "CM-5"], area="iam",
      applies_to=["AWS::IAM::Role"],
      description="PermissionsBoundary is absent. Nothing caps the effective permissions of the role if its policies are later widened.",
      remediation="Attach a Government-approved permissions boundary managed policy via PermissionsBoundary.")
def r_iam_boundary(ctx: TemplateCtx):
    for lid, res in ctx.by_type("AWS::IAM::Role"):
        if not props(res).get("PermissionsBoundary"):
            yield Hit(lid, rtype_of(res), ctx.resource_line(lid, res), "PermissionsBoundary is absent")


@rule(id="CSA-IAM-008", title="Trust or resource policy grants access to Principal *",
      severity=CAT_I, nist=["AC-3", "AC-6", "SC-7"], area="iam",
      applies_to=["AWS::IAM::Role", "AWS::KMS::Key", "AWS::SQS::QueuePolicy", "AWS::SNS::TopicPolicy", "AWS::Lambda::Permission",
                  "AWS::SecretsManager::ResourcePolicy", "AWS::ECR::Repository"],
      overlaps={"CKV_AWS_60", "CKV_AWS_33", "CKV_AWS_51", "F13", "F18", "F21", "F76"}, tags={"exposure"},
      description="An Allow statement (assume-role trust policy, KMS key policy, queue/topic policy or Lambda permission) names Principal * without a restricting Condition such as aws:SourceArn or aws:PrincipalOrgID. Any AWS account can use the resource.",
      remediation="Name the specific principals (account IDs, role ARNs, service principals) or add Condition keys (aws:SourceArn, aws:SourceAccount, aws:PrincipalOrgID).")
def r_public_principal(ctx: TemplateCtx):
    for lid, res in ctx.by_type("AWS::IAM::Role"):
        for s in iter_statements(props(res).get("AssumeRolePolicyDocument")):
            if s.get("Effect") == "Allow" and principal_is_public(s.get("Principal")) and not statement_has_condition(s):
                yield Hit(lid, rtype_of(res), node_line(s, ctx.resource_line(lid, res)), "AssumeRolePolicyDocument allows Principal *")
                break
    for lid, res in ctx.by_type("AWS::KMS::Key", "AWS::SQS::QueuePolicy", "AWS::SNS::TopicPolicy", "AWS::SecretsManager::ResourcePolicy", "AWS::ECR::Repository"):
        p = props(res)
        doc = p.get("KeyPolicy") or p.get("PolicyDocument") or p.get("ResourcePolicy") or p.get("RepositoryPolicyText")
        for s in iter_statements(doc):
            if s.get("Effect") == "Allow" and principal_is_public(s.get("Principal")) and not statement_has_condition(s):
                yield Hit(lid, rtype_of(res), node_line(s, ctx.resource_line(lid, res)), f"policy allows Principal * for {stmt_actions(s)[:3]}")
                break
    for lid, res in ctx.by_type("AWS::Lambda::Permission"):
        p = props(res)
        if p.get("Principal") == "*" and not p.get("SourceArn") and not p.get("SourceAccount") and not p.get("PrincipalOrgID"):
            yield Hit(lid, rtype_of(res), ctx.resource_line(lid, res), "Principal * without SourceArn/SourceAccount")


# ---- Configuration -------------------------------------------------------------------------


@rule(id="CSA-CFG-001", title="EC2 instance does not require IMDSv2 (HttpTokens: required)",
      severity=CAT_II, nist=["CM-6", "AC-6", "SC-7"], area="config", cis=CIS["imdsv2"],
      applies_to=EBS_HOSTS, overlaps={"CKV_AWS_79", "CKV_AWS_341"},
      description="MetadataOptions.HttpTokens is not set to required. IMDSv1 allows an SSRF or local vulnerability to read the instance role credentials without a session token.",
      remediation="Add MetadataOptions: {HttpTokens: required, HttpEndpoint: enabled, HttpPutResponseHopLimit: 1} to the instance, launch template or launch configuration.")
def r_imdsv2(ctx: TemplateCtx):
    for lid, res in ctx.resources.items():
        rtype = rtype_of(res)
        if rtype not in EBS_HOSTS:
            continue
        p = props(res)
        container = p.get("LaunchTemplateData") if rtype == "AWS::EC2::LaunchTemplate" else p
        if not isinstance(container, dict):
            continue
        if rtype == "AWS::EC2::Instance" and p.get("LaunchTemplate"):
            continue  # inherits from the launch template
        mo = container.get("MetadataOptions")
        tokens = mo.get("HttpTokens") if isinstance(mo, dict) else None
        if tokens != "required":
            yield Hit(lid, rtype, ctx.resource_line(lid, res), "MetadataOptions.HttpTokens is absent" if tokens is None else f"HttpTokens={tokens}")


# ---- Secrets and credentials -----------------------------------------------------------------

SECRET_PROPERTY_NAMES = re.compile(r"(masteruserpassword|^password$|adminpassword|dbpassword|secretstring|secretaccesskey|clientsecret|authtoken|^token$|apikey|privatekey)", re.I)
AKIA_RE = re.compile(r"\b(AKIA|ASIA)[0-9A-Z]{16}\b")
USERDATA_SECRET_RE = re.compile(r"(?i)(password|passwd|secret|api[_-]?key|token)\s*[:=]\s*['\"]?[A-Za-z0-9!@#$%^&*()_+\-=]{8,}")
PLACEHOLDER_RE = re.compile(r"(?i)^(\$\{|\{\{|<|change|replace|example|placeholder|xxx|\*+$|password$|changeme|dummy|sample)")


def _param_line(ctx: TemplateCtx, name: str, param: dict) -> int:
    return node_line(param, ctx.find_line(rf'^\s*"?{re.escape(name)}"?\s*:') or 1)


@rule(id="CSA-SEC-001", title="Secret parameter is not marked NoEcho",
      severity=CAT_I, nist=["IA-5", "IA-5(7)", "SC-28"], area="credentials",
      applies_to=["Parameters"], tags={"credential"},
      description="A parameter whose name or description indicates a password, secret, token or private key does not set NoEcho: true. The value is visible in the CloudFormation console, DescribeStacks output and CloudTrail.",
      remediation="Set NoEcho: true on the parameter, or replace it with a Secrets Manager / SSM SecureString dynamic reference.")
def r_param_noecho(ctx: TemplateCtx):
    for name, param in ctx.parameters.items():
        if not parameter_looks_secret(name, param):
            continue
        if not is_true(param.get("NoEcho")):
            yield Hit(name, "Parameter", _param_line(ctx, name, param), f"Parameter {name} (Type {param.get('Type', 'String')}) lacks NoEcho: true")


@rule(id="CSA-SEC-002", title="Secret parameter carries a default value",
      severity=CAT_II, nist=["IA-5", "IA-5(7)"], area="credentials",
      applies_to=["Parameters"], tags={"credential"},
      description="A secret parameter has a Default. The default is stored in the template repository and will be used whenever a deployer does not override it.",
      remediation="Remove the Default and require the value from Secrets Manager or the deployment pipeline.")
def r_param_default_secret(ctx: TemplateCtx):
    for name, param in ctx.parameters.items():
        if parameter_looks_secret(name, param) and param.get("Default") not in (None, ""):
            yield Hit(name, "Parameter", _param_line(ctx, name, param), f"Parameter {name} has Default set")


def walk(node: Any, path: tuple = ()) -> Iterable[tuple[tuple, Any]]:
    if isinstance(node, dict):
        for k, v in node.items():
            yield from walk(v, path + (k,))
    elif isinstance(node, list):
        for i, v in enumerate(node):
            yield from walk(v, path + (i,))
    else:
        yield path, node


@rule(id="CSA-SEC-003", title="Credential appears to be hardcoded in the template",
      severity=CAT_I, nist=["IA-5", "IA-5(7)"], area="credentials",
      applies_to=["*"], overlaps={"CKV_AWS_41", "CKV_AWS_45", "CKV_AWS_46", "CKV_SECRET_2", "CKV_SECRET_6"}, tags={"credential"},
      description="A password/secret/token property, environment variable or user-data line contains a literal value, or an AWS access key ID pattern appears in the template. Anyone with repository access holds the credential.",
      remediation="Move the value to Secrets Manager or SSM SecureString and reference it with a dynamic reference ({{resolve:secretsmanager:...}}) or a NoEcho parameter; rotate the exposed credential.")
def r_hardcoded(ctx: TemplateCtx):
    for lid, res in ctx.resources.items():
        rtype = rtype_of(res)
        for path, value in walk(props(res)):
            if not isinstance(value, str) or len(value) < 6:
                continue
            key = str(path[-1]) if path else ""
            parent_key = str(path[-2]) if len(path) > 1 else ""
            where = ".".join(map(str, path))
            if AKIA_RE.search(value):
                yield Hit(lid, rtype, ctx.resource_line(lid, res), f"AWS access key ID pattern in {where}")
                break
            if SECRET_PROPERTY_NAMES.search(key) and not PLACEHOLDER_RE.match(value) and "{{resolve:" not in value and "${" not in value:
                yield Hit(lid, rtype, ctx.resource_line(lid, res), f"literal value for {where}")
                break
            if parent_key == "Variables" and re.search(r"(?i)secret|password|token|api_?key", key) and "{{resolve:" not in value and not PLACEHOLDER_RE.match(value):
                yield Hit(lid, rtype, ctx.resource_line(lid, res), f"Lambda environment variable {key} holds a literal secret")
                break
            if (key == "UserData" or parent_key == "UserData") and USERDATA_SECRET_RE.search(value) and not re.search(r"(?i)\$\{|\$\(|aws secretsmanager|resolve:", value):
                yield Hit(lid, rtype, ctx.resource_line(lid, res), "UserData assigns a literal password/secret")
                break


@rule(id="CSA-SEC-004", title="Database master password is supplied as a template parameter instead of Secrets Manager",
      severity=CAT_III, nist=["IA-5", "IA-5(7)"], area="credentials",
      applies_to=["AWS::RDS::DBInstance", "AWS::RDS::DBCluster", "AWS::Redshift::Cluster", "AWS::DocDB::DBCluster"], tags={"credential"},
      description="MasterUserPassword is a plain parameter reference. The credential is not generated, stored or rotated by Secrets Manager.",
      remediation="Set ManageMasterUserPassword: true (RDS) or create an AWS::SecretsManager::Secret with GenerateSecretString and reference it with {{resolve:secretsmanager:...}}.")
def r_master_password_param(ctx: TemplateCtx):
    for lid, res in ctx.by_type("AWS::RDS::DBInstance", "AWS::RDS::DBCluster", "AWS::Redshift::Cluster", "AWS::DocDB::DBCluster"):
        p = props(res)
        if any_true(ctx, p.get("ManageMasterUserPassword")):
            continue
        pw = p.get("MasterUserPassword")
        if isinstance(pw, dict) and ctx.param_for(pw):
            yield Hit(lid, rtype_of(res), ctx.prop_line(lid, res, "Properties", "MasterUserPassword"), f"MasterUserPassword: !Ref {ctx.param_for(pw)}")


# ---- Backup, retention, deletion protection ---------------------------------------------------


@rule(id="CSA-BKP-001", title="Stateful resource has no backup or point-in-time recovery configuration",
      severity=CAT_III, nist=["CP-9", "CP-10"], area="recovery",
      applies_to=["AWS::RDS::DBInstance", "AWS::RDS::DBCluster", "AWS::DynamoDB::Table", "AWS::EFS::FileSystem", S3_BUCKET],
      overlaps={"CKV_AWS_28", "CKV_AWS_21", "CKV_AWS_133", "CKV_AWS_139"},
      description="BackupRetentionPeriod is absent or 0 (RDS), PointInTimeRecoverySpecification is not enabled (DynamoDB), BackupPolicy is not ENABLED (EFS) or VersioningConfiguration is not Enabled (S3). Data cannot be restored after corruption or deletion.",
      remediation="Set BackupRetentionPeriod >= 7 with a PreferredBackupWindow; enable PointInTimeRecovery; set BackupPolicy Status ENABLED; enable S3 versioning with lifecycle rules.")
def r_backups(ctx: TemplateCtx):
    for lid, res in ctx.by_type("AWS::RDS::DBInstance", "AWS::RDS::DBCluster"):
        p = props(res)
        if p.get("SourceDBInstanceIdentifier") or p.get("DBClusterIdentifier"):
            continue
        brp = p.get("BackupRetentionPeriod")
        for v in (ctx.resolve(brp) if brp is not None else [None]):
            if v is UNKNOWN:
                continue
            if v is None:
                yield Hit(lid, rtype_of(res), ctx.resource_line(lid, res), "BackupRetentionPeriod is absent (service default 1 day)")
                break
            try:
                zero = int(v) == 0
            except (TypeError, ValueError):
                zero = False
            if zero:
                yield Hit(lid, rtype_of(res), ctx.prop_line(lid, res, "Properties", "BackupRetentionPeriod"), "BackupRetentionPeriod=0 (automated backups disabled)", severity=CAT_II)
                break
    for lid, res in ctx.by_type("AWS::DynamoDB::Table"):
        pitr = props(res).get("PointInTimeRecoverySpecification")
        if not (isinstance(pitr, dict) and any_true(ctx, pitr.get("PointInTimeRecoveryEnabled"))):
            yield Hit(lid, rtype_of(res), ctx.resource_line(lid, res), "PointInTimeRecoverySpecification is not enabled")
    for lid, res in ctx.by_type("AWS::EFS::FileSystem"):
        bp = props(res).get("BackupPolicy")
        if not (isinstance(bp, dict) and bp.get("Status") == "ENABLED"):
            yield Hit(lid, rtype_of(res), ctx.resource_line(lid, res), "BackupPolicy.Status is not ENABLED")
    for lid, res in ctx.by_type(S3_BUCKET):
        vc = props(res).get("VersioningConfiguration")
        if not (isinstance(vc, dict) and vc.get("Status") == "Enabled"):
            yield Hit(lid, rtype_of(res), ctx.resource_line(lid, res), "VersioningConfiguration.Status is not Enabled")


@rule(id="CSA-BKP-002", title="Stateful resource lacks deletion protection or a Retain deletion policy",
      severity=CAT_III, nist=["CP-9", "CM-5"], area="recovery",
      applies_to=["AWS::RDS::DBInstance", "AWS::RDS::DBCluster", "AWS::DynamoDB::Table", S3_BUCKET, "AWS::EFS::FileSystem"],
      overlaps={"CKV_AWS_293", "CKV_AWS_139", "F80"},
      description="DeletionProtection is not true and the resource has no DeletionPolicy of Retain or Snapshot. A stack delete or a template error removes the data.",
      remediation="Set DeletionProtection: true (RDS) / DeletionProtectionEnabled: true (DynamoDB) and DeletionPolicy: Retain or Snapshot with UpdateReplacePolicy on the resource.")
def r_deletion_protection(ctx: TemplateCtx):
    for lid, res in ctx.by_type("AWS::RDS::DBInstance", "AWS::RDS::DBCluster", "AWS::DynamoDB::Table", S3_BUCKET, "AWS::EFS::FileSystem"):
        p = props(res)
        if rtype_of(res) == "AWS::RDS::DBInstance" and (p.get("SourceDBInstanceIdentifier") or p.get("DBClusterIdentifier")):
            continue
        dp = res.get("DeletionPolicy")
        protected = dp in {"Retain", "Snapshot", "RetainExceptOnCreate"}
        if rtype_of(res) in {"AWS::RDS::DBInstance", "AWS::RDS::DBCluster"} and any_true(ctx, p.get("DeletionProtection")):
            protected = True
        if rtype_of(res) == "AWS::DynamoDB::Table" and any_true(ctx, p.get("DeletionProtectionEnabled")):
            protected = True
        if not protected:
            yield Hit(lid, rtype_of(res), ctx.resource_line(lid, res), f"DeletionPolicy={dp}; no deletion protection")


RULES_BY_ID = {r.id: r for r in RULES}


# --------------------------------------------------------------------------------------
# Findings
# --------------------------------------------------------------------------------------


@dataclass
class Finding:
    finding_id: str
    template_path: str
    service_directory: str
    resource_logical_id: str
    resource_type: str
    title: str
    description: str
    evidence: str
    evidence_file: str
    evidence_line: int
    evidence_snippet: str
    severity: str
    severity_justification: str
    nist_controls: list[str]
    cis_aws_v3_id: str | None
    dod_cloud_srg_area: str
    source: str  # custom | checkov | cfn_nag | cfn-lint | custom+checkov ...
    custom_rule_id: str | None
    checkov_ids: list[str]
    cfn_nag_ids: list[str]
    cfn_lint_ids: list[str]
    recommended_remediation: str
    disposition: str
    disposition_note: str
    owner: str
    target_date: str
    tags: list[str]
    risk_score: int = 0
    risk_rank: int = 0
    hit_key: str = ""  # identity of the hit within the resource (ingress rule or policy statement) for rules that report one resource several times

    @property
    def tool_rule_id(self) -> str:
        parts = []
        if self.custom_rule_id:
            parts.append(self.custom_rule_id)
        parts += self.checkov_ids + self.cfn_nag_ids + self.cfn_lint_ids
        return "; ".join(parts)

    @property
    def is_open(self) -> bool:
        return self.disposition == "Open"

    def to_dict(self) -> dict:
        d = self.__dict__.copy()
        d["tool_rule_id"] = self.tool_rule_id
        d["nist_control_families"] = sorted({NIST_CONTROLS.get(c, ("Other", ""))[0] for c in self.nist_controls})
        return d


def make_finding_id(rule_id: str, rel: str, resource: str, discriminator: str = "") -> str:
    """Stable ID from rule, template and resource. ``discriminator`` is the hit key of rules that can report one
    resource several times; it is always applied for such rules, so an ID does not depend on which sibling hits
    exist in a given run."""
    h = hashlib.sha1(f"{rule_id}|{rel}|{resource}".encode()).hexdigest()[:6].upper()
    if discriminator:
        h += "-" + hashlib.sha1(discriminator.encode()).hexdigest()[:4].upper()
    return f"{rule_id}-{h}"


def service_dir_of(rel: str) -> str:
    parts = Path(rel).parts
    return parts[0] if len(parts) > 1 else "(root)"


def finding_from_hit(rule_: Rule, ctx: TemplateCtx, hit: Hit) -> Finding:
    sev = hit.severity or rule_.severity
    line = hit.line or 1
    return Finding(
        finding_id=make_finding_id(rule_.id, ctx.rel, hit.logical_id, hit.key or ""),
        template_path=ctx.rel,
        service_directory=ctx.service_dir,
        resource_logical_id=hit.logical_id,
        resource_type=hit.resource_type,
        title=hit.title or rule_.title,
        description=rule_.description,
        evidence=f"{ctx.rel}:{line} — {hit.detail}",
        evidence_file=ctx.rel,
        evidence_line=line,
        evidence_snippet=ctx.snippet(line),
        severity=sev,
        severity_justification=SEVERITY_JUSTIFICATION[sev],
        nist_controls=list(rule_.nist),
        cis_aws_v3_id=rule_.cis if hit.cis is _UNSET else hit.cis,
        dod_cloud_srg_area=SRG_AREAS[rule_.area],
        source="custom",
        custom_rule_id=rule_.id,
        checkov_ids=[],
        cfn_nag_ids=[],
        cfn_lint_ids=[],
        recommended_remediation=rule_.remediation,
        disposition=hit.disposition or "Open",
        disposition_note="Intended Internet-facing tier; confirm with the system owner and record the acceptance." if hit.disposition else "",
        owner=OWNER_BY_AREA[rule_.area],
        target_date=TARGET_DATE_PLACEHOLDER[sev],
        tags=sorted(rule_.tags),
        hit_key=hit.key or "",
    )


# --------------------------------------------------------------------------------------
# External tools
# --------------------------------------------------------------------------------------


@dataclass
class ToolResult:
    name: str
    version: str
    status: str  # ran | skipped | failed
    note: str = ""
    raw_count: int = 0
    duration_s: float = 0.0
    unparsed: list[str] = field(default_factory=list)   # templates the tool should have covered but did not
    excluded: list[str] = field(default_factory=list)   # templates deliberately out of the tool's scope


def run(cmd: list[str], cwd: Path, timeout: int = 3600) -> subprocess.CompletedProcess:
    return subprocess.run(cmd, cwd=str(cwd), capture_output=True, text=True, timeout=timeout, check=False)


def tool_version(cmd: list[str]) -> str:
    try:
        cp = subprocess.run(cmd, capture_output=True, text=True, timeout=120, check=False)
        out = (cp.stdout or cp.stderr).strip().splitlines()
        return out[0].strip() if out else "unknown"
    except (OSError, subprocess.TimeoutExpired):
        return "unavailable"


def find_cfn_nag() -> str | None:
    exe = shutil.which("cfn_nag_scan")
    if exe:
        return exe
    candidates: list[str] = []
    home = Path.home()
    for pattern in ("~/.local/share/gem/ruby/*/bin/cfn_nag_scan", "~/.gem/ruby/*/bin/cfn_nag_scan", "/var/lib/gems/*/bin/cfn_nag_scan",
                    "/usr/local/bundle/bin/cfn_nag_scan", "~/.rbenv/shims/cfn_nag_scan"):
        candidates += [str(p) for p in Path("/").glob(os.path.expanduser(pattern).lstrip("/"))]
    try:
        cp = subprocess.run(["gem", "environment", "gempath"], capture_output=True, text=True, timeout=60, check=False)
        for gp in cp.stdout.strip().split(":"):
            c = Path(gp) / "bin" / "cfn_nag_scan"
            if c.exists():
                candidates.append(str(c))
    except (OSError, subprocess.TimeoutExpired):
        pass
    _ = home
    return candidates[0] if candidates else None


# Checkov / cfn_nag control mapping for findings that no custom rule covers. Keyword rules are applied
# in order; the first match wins. (nist, area, severity, cis)
KEYWORD_MAP: list[tuple[str, tuple[list[str], str, str, str | None]]] = [
    (r"0\.0\.0\.0|::/0|open to the world|public(ly)? accessib|publicly|world|internet", (["SC-7", "SC-7(5)"], "boundary", CAT_II, None)),
    (r"imdsv2|metadata", (["CM-6", "AC-6"], "config", CAT_II, CIS["imdsv2"])),
    (r"encrypt|kms|cmk|customer managed key|ssl|tls|https|secure transport|in transit|in-transit", (["SC-28", "SC-28(1)", "SC-8", "SC-8(1)"], "encryption_rest", CAT_II, None)),
    (r"log|trail|audit|flow", (["AU-2", "AU-12"], "logging", CAT_III, None)),
    (r"monitor|tracing|x-ray|alarm|insights|enhanced", (["SI-4"], "monitoring", CAT_III, None)),
    (r"passrole|wildcard|\*|admin|privilege|iam|policy|permission|principal|assume|trust|access key|mfa|credential|password|secret", (["AC-6", "AC-3", "IA-5"], "iam", CAT_II, None)),
    (r"backup|retention|point.in.time|versioning|snapshot|deletion protection|termination protection|recovery|multi-az|multi az", (["CP-9", "CP-10"], "recovery", CAT_III, None)),
    (r"security group|ingress|egress|port|cidr|nacl|subnet|public ip|endpoint|vpc", (["SC-7"], "boundary", CAT_II, None)),
    (r"deprecated|runtime|version|latest|upgrade|patch|minor", (["SI-2", "CM-6"], "config", CAT_III, None)),
]

# Checkov IDs whose default severity in the keyword map is too high/low for this baseline.
CHECKOV_SEVERITY_OVERRIDE = {
    "CKV_AWS_18": CAT_III, "CKV_AWS_21": CAT_III, "CKV_AWS_144": CAT_III, "CKV_AWS_145": CAT_III, "CKV_AWS_19": CAT_II,
    "CKV_AWS_116": CAT_III, "CKV_AWS_115": CAT_III, "CKV_AWS_173": CAT_III, "CKV_AWS_158": CAT_III, "CKV_AWS_66": CAT_III,
    "CKV_AWS_7": CAT_III, "CKV_AWS_33": CAT_II, "CKV_AWS_23": CAT_III, "CKV_AWS_100": CAT_III, "CKV_AWS_1": CAT_II,
    "CKV_AWS_107": CAT_II, "CKV_AWS_109": CAT_II, "CKV_AWS_110": CAT_II, "CKV_AWS_111": CAT_II, "CKV_AWS_108": CAT_II,
    "CKV_AWS_24": CAT_II, "CKV_AWS_25": CAT_II, "CKV_AWS_260": CAT_III, "CKV_AWS_17": CAT_II, "CKV_AWS_20": CAT_II,
    "CKV_AWS_57": CAT_II, "CKV_AWS_41": CAT_II, "CKV_AWS_45": CAT_II, "CKV_AWS_46": CAT_II, "CKV_AWS_16": CAT_II,
    "CKV_AWS_60": CAT_II, "CKV_AWS_61": CAT_II, "CKV_AWS_40": CAT_II, "CKV_AWS_2": CAT_II, "CKV_AWS_103": CAT_II,
    "CKV_AWS_126": CAT_III, "CKV_AWS_135": CAT_III, "CKV_AWS_88": CAT_II, "CKV_AWS_59": CAT_II, "CKV_AWS_53": CAT_II,
    "CKV_AWS_54": CAT_II, "CKV_AWS_55": CAT_II, "CKV_AWS_56": CAT_II, "CKV_AWS_79": CAT_II, "CKV_AWS_38": CAT_II,
    "CKV_AWS_39": CAT_II, "CKV_AWS_58": CAT_II, "CKV_AWS_35": CAT_II, "CKV_AWS_36": CAT_III, "CKV_AWS_67": CAT_III,
    "CKV_AWS_252": CAT_III, "CKV_AWS_50": CAT_III, "CKV_AWS_117": CAT_III, "CKV_AWS_86": CAT_III, "CKV_AWS_68": CAT_III,
    "CKV_AWS_34": CAT_II, "CKV_AWS_174": CAT_II, "CKV_AWS_28": CAT_III, "CKV_AWS_119": CAT_III, "CKV_AWS_27": CAT_III,
    "CKV_AWS_26": CAT_II, "CKV_AWS_42": CAT_II, "CKV_AWS_3": CAT_II, "CKV_AWS_8": CAT_II, "CKV_AWS_189": CAT_II,
    "CKV_AWS_64": CAT_II, "CKV_AWS_74": CAT_II, "CKV_AWS_96": CAT_II, "CKV_AWS_44": CAT_II, "CKV_AWS_4": CAT_III,
    "CKV_AWS_5": CAT_II, "CKV_AWS_29": CAT_II, "CKV_AWS_30": CAT_II, "CKV_AWS_31": CAT_II, "CKV_AWS_37": CAT_II,
    "CKV_AWS_65": CAT_III, "CKV_AWS_69": CAT_II, "CKV_AWS_83": CAT_II, "CKV_AWS_84": CAT_III, "CKV_AWS_90": CAT_III,
    "CKV_AWS_91": CAT_III, "CKV_AWS_92": CAT_III, "CKV_AWS_93": CAT_II, "CKV_AWS_95": CAT_III, "CKV_AWS_97": CAT_III,
    "CKV_AWS_105": CAT_III, "CKV_AWS_118": CAT_III, "CKV_AWS_120": CAT_III, "CKV_AWS_129": CAT_III, "CKV_AWS_130": CAT_III,
    "CKV_AWS_131": CAT_III, "CKV_AWS_150": CAT_III, "CKV_AWS_157": CAT_III, "CKV_AWS_161": CAT_III, "CKV_AWS_162": CAT_III,
    "CKV_AWS_163": CAT_III, "CKV_AWS_166": CAT_III, "CKV_AWS_186": CAT_III, "CKV_AWS_209": CAT_III, "CKV_AWS_226": CAT_III,
    "CKV_AWS_240": CAT_II, "CKV_AWS_241": CAT_II, "CKV_AWS_293": CAT_III, "CKV_AWS_338": CAT_III, "CKV_AWS_339": CAT_III,
    "CKV_AWS_341": CAT_II, "CKV_AWS_363": CAT_III, "CKV_AWS_378": CAT_II, "CKV_AWS_382": CAT_III,
    "CKV_AWS_172": CAT_III, "CKV_AWS_51": CAT_II, "CKV_AWS_136": CAT_III,
}
CHECKOV_CIS = {"CKV_AWS_24": CIS["sg_admin_ipv4"], "CKV_AWS_25": CIS["sg_admin_ipv4"], "CKV_AWS_79": CIS["imdsv2"], "CKV_AWS_341": CIS["imdsv2"],
               "CKV_AWS_16": CIS["rds_encryption"], "CKV_AWS_17": CIS["rds_public"], "CKV_AWS_3": CIS["ebs_encryption"], "CKV_AWS_8": CIS["ebs_encryption"],
               "CKV_AWS_42": CIS["efs_encryption"], "CKV_AWS_35": CIS["cloudtrail_kms"], "CKV_AWS_36": CIS["cloudtrail_validation"],
               "CKV_AWS_67": CIS["cloudtrail_multiregion"], "CKV_AWS_7": CIS["kms_rotation"], "CKV_AWS_53": CIS["s3_block_public"],
               "CKV_AWS_54": CIS["s3_block_public"], "CKV_AWS_55": CIS["s3_block_public"], "CKV_AWS_56": CIS["s3_block_public"],
               "CKV_AWS_226": CIS["rds_minor_upgrade"], "CKV_AWS_378": None}
CFN_NAG_SEVERITY_OVERRIDE = {"F1000": CAT_III, "W28": CAT_III, "W35": CAT_III, "W51": CAT_III, "W58": CAT_III, "W89": CAT_III, "W92": CAT_III,
                             "W41": CAT_II, "W47": CAT_II, "W48": CAT_III, "W84": CAT_III, "W86": CAT_III, "W5": CAT_III, "W29": CAT_III,
                             "W36": CAT_III, "W42": CAT_II, "W2": CAT_II, "W9": CAT_III, "W40": CAT_III, "W11": CAT_II, "W12": CAT_II,
                             "W13": CAT_II, "W76": CAT_III, "W77": CAT_III, "W68": CAT_III, "W69": CAT_III, "W64": CAT_III, "W74": CAT_III,
                             "W78": CAT_III, "W75": CAT_III, "W57": CAT_III, "W59": CAT_II, "W56": CAT_II, "W26": CAT_III, "W72": CAT_III}


# cfn_nag checks that report the same weakness as a Checkov check; when both fire on one resource the cfn_nag
# ID is attached to the Checkov finding instead of producing a second row.
CFN_NAG_EQUIVALENT_CHECKOV = {"W36": "CKV_AWS_23", "W89": "CKV_AWS_117", "W55": "CKV_AWS_103", "W10": "CKV_AWS_86", "W26": "CKV_AWS_92",
                              "W52": "CKV_AWS_91", "W74": "CKV_AWS_119", "F27": "CKV_AWS_16", "F1": "CKV_AWS_3", "W59": "CKV_AWS_59",
                              "W69": "CKV_AWS_76", "W35": "CKV_AWS_18", "W41": "CKV_AWS_19", "W47": "CKV_AWS_26", "W48": "CKV_AWS_27",
                              "W84": "CKV_AWS_158", "W86": "CKV_AWS_66", "W70": "CKV_AWS_174", "F22": "CKV_AWS_17", "F80": "CKV_AWS_293"}


def classify_external(check_id: str, name: str, source: str) -> tuple[list[str], str, str, str | None, str | None]:
    """Return (nist, area, severity, cis, overlapping custom rule id) for a Checkov / cfn_nag check."""
    if source == "cfn_nag" and check_id == "FATAL":
        return ["CM-2", "CM-6"], "validity", CAT_III, None, None
    overlap: Rule | None = next((r for r in RULES if check_id in r.overlaps), None)
    if overlap is not None:
        # The custom rule that covers this weakness did not fire on the resource (otherwise the hit would have been
        # merged into it), so the tool reports a broader condition than the rule pack's CAT I criteria. Keep the
        # rule's control mapping but cap the severity at CAT II pending analyst review.
        nist, area, sev, cis = list(overlap.nist), overlap.area, (CAT_II if overlap.severity == CAT_I else overlap.severity), overlap.cis
    else:
        nist, area, sev, cis = ["CM-6"], "config", CAT_III, None
        for pattern, mapping in KEYWORD_MAP:
            if re.search(pattern, name, re.I):
                nist, area, sev, cis = mapping
                break
    if source == "cfn_nag":
        sev = CFN_NAG_SEVERITY_OVERRIDE.get(check_id, CAT_II if check_id.startswith("F") else sev)
    else:
        sev = CHECKOV_SEVERITY_OVERRIDE.get(check_id, sev)
        cis = CHECKOV_CIS.get(check_id, cis)
    # CAT I is reserved for conditions the custom rule pack verified on the parsed template (tool hits that
    # confirm them are merged into that finding). Tool-only hits rely on pattern heuristics and are capped at CAT II.
    if sev == CAT_I:
        sev = CAT_II
    return nist, area, sev, cis, overlap.id if overlap else None


def resource_line_in_file(lines: list[str], logical_id: str) -> int:
    rx = re.compile(rf'^\s*"?{re.escape(logical_id)}"?\s*:')
    for i, text in enumerate(lines):
        if rx.search(text):
            return i + 1
    return 1


@dataclass
class ExternalHit:
    source: str
    check_id: str
    name: str
    rel: str
    logical_id: str
    resource_type: str
    line: int
    guideline: str = ""
    message: str = ""


def parse_json_prefix(out: str, opener: str = "[{") -> object | None:
    """Parse the first JSON document in tool output that may be preceded by log lines."""
    out = out.strip()
    starts = [i for i in (out.find(c) for c in opener) if i >= 0]
    if not starts:
        return None
    try:
        return json.loads(out[min(starts):])
    except json.JSONDecodeError:
        return None


def has_rain_directive(path: Path) -> bool:
    try:
        return "!Rain::" in path.read_text(encoding="utf-8", errors="replace")
    except OSError:
        return False


def checkov_parsing_errors(data: object) -> int:
    """Number of files Checkov failed to parse. Checkov exits 0 for a malformed template and reports it only in
    ``summary.parsing_errors`` (or ``results.parsing_errors``), so that counter, not the exit code, decides coverage."""
    n = 0
    for rep in data if isinstance(data, list) else [data]:
        if not isinstance(rep, dict):
            continue
        summary = rep.get("summary") if isinstance(rep.get("summary"), dict) else rep
        pe = summary.get("parsing_errors", 0)
        n += len(pe) if isinstance(pe, list) else int(pe or 0)
        pe = (rep.get("results") or {}).get("parsing_errors") if isinstance(rep.get("results"), dict) else None
        n += len(pe) if isinstance(pe, list) else 0
    return n


def checkov_hits_from_report(data: object) -> list[ExternalHit]:
    hits: list[ExternalHit] = []
    reports = data if isinstance(data, list) else [data]
    for rep in reports:
        if not isinstance(rep, dict):
            continue
        framework = rep.get("check_type", "")
        for f in (rep.get("results") or {}).get("failed_checks") or []:
            rel = str(f.get("file_path", "")).lstrip("/")
            rel = str(Path(rel)) if rel else rel
            resource = str(f.get("resource", ""))
            rtype, _, lid = resource.partition(".")
            if framework == "terraform" or framework == "secrets":
                rtype, lid = framework + ":" + rtype, lid or resource
            rng = f.get("file_line_range") or [1, 1]
            check = f.get("check")
            name = check.get("name", "") if isinstance(check, dict) else str(f.get("check_name", ""))
            hits.append(ExternalHit("checkov", str(f.get("check_id", "")), name, rel, lid or resource, rtype,
                                    int(rng[0]) if rng else 1, str(f.get("guideline") or "")))
    return hits


def run_checkov(repo: Path, templates: list[Path], terraform: list[Path], tr: ToolResult) -> list[ExternalHit]:
    """Run Checkov once per template so that a template Checkov cannot parse (for example a
    Rain module reference used as a resource ``Type``) does not abort the whole scan."""
    exe = shutil.which("checkov")
    if not exe:
        tr.status, tr.note = "skipped", "checkov executable not found (pip install checkov)"
        return []
    tr.version = tool_version([exe, "--version"])
    t0 = time.time()
    base = [exe, "-o", "json", "--quiet", "--compact"]

    def scan_file(t: Path) -> tuple[Path, list[ExternalHit] | None]:
        try:
            cp = run(base + ["-f", str(t.relative_to(repo)), "--framework", "cloudformation", "secrets"], repo, timeout=600)
        except subprocess.TimeoutExpired:
            return t, None
        data = parse_json_prefix(cp.stdout)
        if data is None or checkov_parsing_errors(data):
            return t, None
        hits = checkov_hits_from_report(data)
        for h in hits:
            h.rel = str(t.relative_to(repo))
        return t, hits

    hits: list[ExternalHit] = []
    failed: list[str] = []
    with concurrent.futures.ThreadPoolExecutor(max_workers=max(2, min(8, os.cpu_count() or 2))) as pool:
        for t, result in pool.map(scan_file, templates):
            if result is None:
                failed.append(str(t.relative_to(repo)))
            else:
                hits += result
    for tf_dir in sorted({t.parent for t in terraform}):
        try:
            cp = run(base + ["-d", str(tf_dir.relative_to(repo)), "--framework", "terraform"], repo, timeout=1200)
        except subprocess.TimeoutExpired:
            failed.append(str(tf_dir.relative_to(repo)))
            continue
        data = parse_json_prefix(cp.stdout)
        if data is None or checkov_parsing_errors(data):
            failed.append(str(tf_dir.relative_to(repo)))
            continue
        for h in checkov_hits_from_report(data):
            h.rel = str((tf_dir / Path(h.rel).name).relative_to(repo)) if not h.rel.startswith(str(tf_dir.relative_to(repo))) else h.rel
            hits.append(h)
    tr.duration_s = round(time.time() - t0, 1)
    tr.status, tr.raw_count = "ran", len(hits)
    tr.unparsed = failed
    if failed:
        tr.note = f"{len(failed)} template(s) could not be parsed by checkov and were covered by the custom rule pack only: " + ", ".join(failed)
    return hits


def run_cfn_nag(repo: Path, templates: list[Path], tr: ToolResult) -> list[ExternalHit]:
    """Run cfn_nag once per template (a single template that cfn_nag cannot parse would otherwise fail the aggregate scan)."""
    exe = find_cfn_nag()
    if not exe:
        tr.status, tr.note = "skipped", "cfn_nag_scan not found on PATH or in RubyGems bin directories (gem install cfn-nag)"
        return []
    tr.version = tool_version([exe, "--version"])
    t0 = time.time()

    def scan_file(t: Path) -> tuple[Path, list[ExternalHit] | None]:
        try:
            cp = run([exe, "--input-path", str(t.relative_to(repo)), "--output-format", "json"], repo, timeout=600)
        except subprocess.TimeoutExpired:
            return t, None
        data = parse_json_prefix(cp.stdout, "[")
        if not isinstance(data, list):
            return t, None
        rel = str(t.relative_to(repo))
        out: list[ExternalHit] = []
        for entry in data:
            violations = (entry.get("file_results") or {}).get("violations") or []
            if any(v.get("id") == "FATAL" for v in violations):
                return t, None  # cfn_nag could not parse the template; it reports that as a FATAL pseudo-violation
            for v in violations:
                ids = v.get("logical_resource_ids") or [""]
                lines = v.get("line_numbers") or []
                for i, lid in enumerate(ids):
                    line = int(lines[i]) if i < len(lines) and lines[i] else 1
                    out.append(ExternalHit("cfn_nag", str(v.get("id", "")), str(v.get("message", "")), rel, str(lid), "", line))
        return t, out

    hits: list[ExternalHit] = []
    failed: list[str] = []
    with concurrent.futures.ThreadPoolExecutor(max_workers=max(2, min(8, os.cpu_count() or 2))) as pool:
        for t, result in pool.map(scan_file, templates):
            if result is None:
                failed.append(str(t.relative_to(repo)))
            else:
                hits += result
    tr.duration_s = round(time.time() - t0, 1)
    tr.status, tr.raw_count = "ran", len(hits)
    tr.unparsed = failed
    if failed:
        tr.note = f"{len(failed)} template(s) could not be parsed by cfn_nag: " + ", ".join(failed)
    return hits


# cfn-lint exit codes that carry a result list: 0 clean, 2 errors, 4 warnings, 8 informational and their sums.
# 1 is an invocation error (bad arguments, unreadable input) and produces no result list.
LINT_RESULT_CODES = {0, 2, 4, 6, 8, 10, 12, 14}


def run_cfn_lint(repo: Path, templates: list[Path], tr: ToolResult) -> tuple[list[ExternalHit], int]:
    """Run cfn-lint the way the repository's own lint gate does (scripts/lint-single.sh): macro examples are
    not linted, templates that use ``!Rain::`` directives are packaged with ``rain pkg`` first when ``rain`` is
    available and skipped otherwise, and module fragments under ``RainModules/`` are not standalone templates."""
    exe = shutil.which("cfn-lint")
    if not exe:
        tr.status, tr.note = "skipped", "cfn-lint executable not found (pip install cfn-lint)"
        return [], 0
    tr.version = tool_version([exe, "--version"])
    tr.note = ""
    t0 = time.time()
    hits: list[ExternalHit] = []
    warnings = 0
    rain = shutil.which("rain")
    plain: list[Path] = []
    rain_templates: list[Path] = []
    excluded: list[str] = []
    failed: list[str] = []
    for t in templates:
        rel = t.relative_to(repo).as_posix()
        if "MacrosExamples/" in rel or rel.startswith("RainModules/"):
            excluded.append(rel)
        elif has_rain_directive(t):
            rain_templates.append(t)
        else:
            plain.append(t)
    base = [exe, "--format", "json", "--non-zero-exit-code", "error"]
    if (repo / ".cfnlintrc").exists():
        base += ["--config-file", str(repo / ".cfnlintrc")]

    def collect(data: list, filename: str | None = None) -> None:
        nonlocal warnings
        for m in data:
            if not isinstance(m, dict):
                continue
            if m.get("Level", "") != "Error":
                warnings += 1
                continue
            rule_ = m.get("Rule") or {}
            loc = m.get("Location") or {}
            path = loc.get("Path") or []
            lid = str(path[1]) if len(path) > 1 and path[0] == "Resources" else (str(path[0]) if path else "")
            rel = filename or Path(str(m.get("Filename", ""))).as_posix().lstrip("./")
            hits.append(ExternalHit("cfn-lint", str(rule_.get("Id", "")), str(rule_.get("ShortDescription", "")), rel, lid, "",
                                    int(((loc.get("Start") or {}).get("LineNumber")) or 1), message=str(m.get("Message", ""))))

    batch = 40
    for i in range(0, len(plain), batch):
        chunk = [str(t.relative_to(repo)) for t in plain[i:i + batch]]
        try:
            cp = run(base + ["--"] + chunk, repo, timeout=1800)
        except subprocess.TimeoutExpired:
            tr.note += f" batch {i // batch + 1} timed out;"
            failed += [Path(c).as_posix() for c in chunk]
            continue
        data = parse_json_prefix(cp.stdout, "[")
        if cp.returncode not in LINT_RESULT_CODES or not isinstance(data, list):
            # No result list means the invocation itself failed; none of the chunk was linted.
            tr.note += f" batch {i // batch + 1} returned exit code {cp.returncode} without a JSON result list;"
            failed += [Path(c).as_posix() for c in chunk]
            continue
        invocation_errors = [m.get("Message", "") for m in data if isinstance(m, dict) and not m.get("Filename")]
        if invocation_errors:
            # cfn-lint stops before linting any file of the invocation when one argument cannot be processed (E0003 etc.).
            tr.note += f" batch {i // batch + 1} aborted: {str(invocation_errors[0])[:120]};"
            failed += [Path(c).as_posix() for c in chunk]
            continue
        collect(data)
    for t in rain_templates:
        rel = t.relative_to(repo).as_posix()
        if not rain:
            failed.append(rel)
            continue
        try:
            pkg = run([rain, "pkg", "-x", rel], repo, timeout=600)
            if pkg.returncode != 0:
                failed.append(rel)
                continue
            cp = subprocess.run(base, cwd=str(repo), input=pkg.stdout, capture_output=True, text=True, timeout=600, check=False)
        except subprocess.TimeoutExpired:
            failed.append(rel)
            continue
        data = parse_json_prefix(cp.stdout, "[")
        if cp.returncode not in LINT_RESULT_CODES or not isinstance(data, list):
            failed.append(rel)
            continue
        collect(data, rel)
    tr.duration_s = round(time.time() - t0, 1)
    tr.status, tr.raw_count = "ran", len(hits)
    tr.unparsed = failed
    tr.excluded = excluded
    note = f"{warnings} warning/informational messages not treated as findings"
    if excluded:
        note += (f"; {len(excluded)} template(s) not linted per the repository lint convention (macro examples and Rain module fragments): "
                 + ", ".join(excluded))
    if failed:
        note += (f"; {len(failed)} template(s) could not be linted (invocation failed or timed out, or `rain pkg` unavailable or failed): "
                 + ", ".join(failed))
    tr.note = (tr.note + " " + note).strip()
    return hits, warnings


def normalize_external(hit: ExternalHit, repo: Path, resource_types: dict[tuple[str, str], str]) -> Finding:
    rtype = hit.resource_type or resource_types.get((hit.rel, hit.logical_id), "")
    if hit.source == "cfn-lint":
        nist, area, sev, cis, overlap = ["CM-2", "CM-6"], "validity", CAT_III, None, None
        title = f"Template does not pass cfn-lint ({hit.check_id})"
        desc = ("cfn-lint reports an error-level rule violation. A template that does not validate cannot be relied on as a "
                "configuration baseline and may fail or deploy unintended settings. Message: " + hit.message)
        rem = "Correct the template so `cfn-lint` passes with the repository .cfnlintrc; add the check to the CI gate."
    else:
        nist, area, sev, cis, overlap = classify_external(hit.check_id, hit.name, hit.source)
        title = f"{hit.name} ({hit.source} {hit.check_id})" if hit.name else f"{hit.source} {hit.check_id}"
        desc = f"{hit.source} check {hit.check_id} failed for this resource: {hit.name}."
        if hit.guideline:
            desc += f" Guideline: {hit.guideline}"
        rem = RULES_BY_ID[overlap].remediation if overlap else "Apply the secure setting described by the check and re-run the assessment."
    snippet = ""
    p = repo / hit.rel
    if p.exists():
        try:
            lines = p.read_text(encoding="utf-8", errors="replace").splitlines()
            if 1 <= hit.line <= len(lines):
                snippet = lines[hit.line - 1].strip()[:140]
        except OSError:
            pass
    tool_key = f"{hit.source}:{hit.check_id}"
    f = Finding(
        finding_id=make_finding_id(hit.check_id.replace("_", "-"), hit.rel, hit.logical_id),
        template_path=hit.rel, service_directory=service_dir_of(hit.rel), resource_logical_id=hit.logical_id, resource_type=rtype,
        title=title, description=desc, evidence=f"{hit.rel}:{hit.line} — {hit.name or hit.message}", evidence_file=hit.rel, evidence_line=hit.line,
        evidence_snippet=snippet, severity=sev, severity_justification=SEVERITY_JUSTIFICATION[sev], nist_controls=nist, cis_aws_v3_id=cis,
        dod_cloud_srg_area=SRG_AREAS[area], source=hit.source, custom_rule_id=None,
        checkov_ids=[hit.check_id] if hit.source == "checkov" else [], cfn_nag_ids=[hit.check_id] if hit.source == "cfn_nag" else [],
        cfn_lint_ids=[hit.check_id] if hit.source == "cfn-lint" else [], recommended_remediation=rem, disposition="Open", disposition_note="",
        owner=OWNER_BY_AREA[area], target_date=TARGET_DATE_PLACEHOLDER[sev], tags=sorted(RULES_BY_ID[overlap].tags) if overlap else [],
    )
    _ = tool_key
    return f


# Scanner checks that name one port; they attach only to the custom ingress finding whose port range covers it.
PORT_SPECIFIC_CHECKS = {"CKV_AWS_24": 22, "CKV_AWS_25": 3389, "CKV_AWS_260": 80}


def ingress_key_covers_port(key: str, port: int) -> bool:
    """True when an ``ingress_key`` (``cidr|from-to/proto``) spans ``port`` (a -1 or None bound means all ports)."""
    m = re.fullmatch(r"[^|]*\|(-?\d+|None)-(-?\d+|None)/(.*)", key)
    if not m:
        return False
    lo, hi, proto = m.group(1), m.group(2), m.group(3)
    if proto not in ("tcp", "-1", "all"):
        return False
    lo_i = -1 if lo == "None" else int(lo)
    hi_i = -1 if hi == "None" else int(hi)
    return lo_i == -1 or hi_i == -1 or lo_i <= port <= hi_i


def merge_external(custom: list[Finding], external: list[ExternalHit], repo: Path, resource_types: dict[tuple[str, str], str],
                   assessed: set[str]) -> list[Finding]:
    """Attach tool IDs to custom findings that cover the same weakness on the same resource; keep the rest as tool findings.

    Hits on files outside ``assessed`` (generated JSON twins, non-template files) are dropped.
    """
    # Every custom finding for (template, resource, rule); rules that report one resource several times (Hit.key)
    # contribute several candidates and an external hit attaches only to the candidate on its own line, so tool
    # evidence is never credited to a different ingress rule or policy statement.
    by_key: dict[tuple[str, str, str], list[Finding]] = {}
    for f in custom:
        if f.custom_rule_id:
            by_key.setdefault((f.template_path, f.resource_logical_id, f.custom_rule_id), []).append(f)

    def candidate(hit: ExternalHit, rule_id: str) -> Finding | None:
        cands = by_key.get((hit.rel, hit.logical_id, rule_id), [])
        if len(cands) == 1:
            return cands[0]
        port = PORT_SPECIFIC_CHECKS.get(hit.check_id)
        if port is not None:
            cands = [c for c in cands if ingress_key_covers_port(c.hit_key, port)]
            if len(cands) == 1:
                return cands[0]
        same_line = [c for c in cands if c.evidence_line == hit.line]
        return same_line[0] if len(same_line) == 1 else None
    extra: list[Finding] = []
    tool_only: dict[tuple[str, str, str], Finding] = {}
    seen: set[tuple[str, str, str]] = set()
    # Checkov first so that cfn_nag hits can attach to an equivalent Checkov finding.
    for hit in sorted(external, key=lambda h: h.source != "checkov"):
        if hit.rel not in assessed:
            continue
        dedupe = (hit.source, hit.check_id, f"{hit.rel}|{hit.logical_id}")
        if dedupe in seen:
            continue
        seen.add(dedupe)
        merged = False
        for r in RULES:
            if hit.check_id in r.overlaps:
                f = candidate(hit, r.id)
                if f is not None:
                    (f.checkov_ids if hit.source == "checkov" else f.cfn_nag_ids).append(hit.check_id)
                    if hit.source not in f.source:
                        f.source += "+" + hit.source
                    merged = True
                    break
        if merged:
            continue
        twin = CFN_NAG_EQUIVALENT_CHECKOV.get(hit.check_id) if hit.source == "cfn_nag" else None
        if twin is not None:
            f = tool_only.get((hit.rel, hit.logical_id, twin))
            if f is not None:
                f.cfn_nag_ids.append(hit.check_id)
                f.source = "checkov+cfn_nag"
                continue
        f = normalize_external(hit, repo, resource_types)
        tool_only[(hit.rel, hit.logical_id, hit.check_id)] = f
        extra.append(f)
    return custom + extra


# --------------------------------------------------------------------------------------
# Risk ranking, baseline, dispositions
# --------------------------------------------------------------------------------------


def score(f: Finding) -> int:
    s = SEVERITY_BASE_SCORE[f.severity]
    if "exposure" in f.tags:
        s += 40
    if "credential" in f.tags:
        s += 30
    if "datastore" in f.tags:
        s += 20
    s += 10 * (len(f.checkov_ids) > 0) + 10 * (len(f.cfn_nag_ids) > 0)
    if f.custom_rule_id:
        s += 5
    if f.service_directory.lower() in {"solutions", "vpc", "rds", "s3", "ec2", "iam"}:
        s += 3
    return s


def rank(findings: list[Finding]) -> None:
    for f in findings:
        f.risk_score = score(f)
    open_first = sorted(findings, key=lambda f: (not f.is_open, -f.risk_score, SEVERITY_ORDER[f.severity], f.template_path, f.resource_logical_id, f.finding_id))
    for i, f in enumerate(open_first, start=1):
        f.risk_rank = i
    findings.sort(key=lambda f: f.risk_rank)


def path_in(path: str, unprocessed: Iterable[str]) -> bool:
    """True when ``path`` is one of ``unprocessed`` or lies under a directory listed there (Checkov scans Terraform per directory)."""
    return any(path == u or path.startswith(u.rstrip("/") + "/") for u in unprocessed)


def coverage_gap(f: Finding, tools: list[ToolResult], parsed: set[str], rule_failures: dict[str, list[str]] | None = None) -> str:
    """Why the current run could not re-evaluate ``f``; empty when every source that produced it was rerun on its file."""
    by_name = {t.name: t for t in tools}
    for src in f.source.split("+"):
        if src == "custom":
            if f.template_path not in parsed:
                return "its template was not parsed by the custom rule pack in this run"
            if f.custom_rule_id and f.custom_rule_id in (rule_failures or {}).get(f.template_path, []):
                return f"rule {f.custom_rule_id} failed on the template in this run"
            continue
        t = by_name.get(src)
        if t is None or t.status != "ran":
            reason = t.note if t and t.note else "not configured"
            return f"{src} did not run in this run ({reason})"
        if path_in(f.template_path, t.unparsed):
            return f"{src} could not process the file in this run"
    return ""


def apply_baseline(findings: list[Finding], baseline_path: Path | None, dispositions_path: Path | None,
                   tools: list[ToolResult] | None = None, parsed: set[str] | None = None,
                   rule_failures: dict[str, list[str]] | None = None) -> tuple[list[Finding], dict | None]:
    """Carry baseline findings that are absent from the current run.

    A missing finding is ``Remediated in PR`` only when every tool that produced it ran again on
    the same template. When a scanner was skipped, failed or could not parse that template the
    finding keeps its previous disposition and its note records that it was not re-evaluated.
    """
    before: dict | None = None
    current_ids = {f.finding_id for f in findings}
    if baseline_path and baseline_path.exists():
        base = json.loads(baseline_path.read_text())
        before = base.get("summary")
        base_sha = str(base.get("metadata", {}).get("commit_sha", "?"))[:12]
        for d in base.get("findings", []):
            if d["finding_id"] in current_ids:
                continue
            if d.get("disposition") in {"Not applicable"}:
                continue
            carried = Finding(**{k: v for k, v in d.items() if k in {fl.name for fl in fields(Finding)}})
            gap = coverage_gap(carried, tools or [], parsed if parsed is not None else {carried.template_path}, rule_failures)
            if gap:
                carried.disposition_note = f"Not re-evaluated: {gap}; disposition carried forward from baseline commit {base_sha}."
            else:
                carried.disposition = "Remediated in PR"
                carried.disposition_note = f"Finding no longer detected after remediation (baseline commit {base_sha})."
            findings.append(carried)
    if dispositions_path and dispositions_path.exists():
        overrides = json.loads(dispositions_path.read_text())
        for f in findings:
            o = overrides.get(f.finding_id)
            if isinstance(o, dict) and o.get("disposition") in DISPOSITIONS:
                f.disposition = o["disposition"]
                f.disposition_note = o.get("note", f.disposition_note)
    return findings, before


# --------------------------------------------------------------------------------------
# Discovery and assessment
# --------------------------------------------------------------------------------------


def discover_templates(repo: Path, include_generated_json: bool = False) -> tuple[list[Path], list[Path], list[str], dict[Path, Path], list[str]]:
    """Return (templates, terraform files, skipped non-template files, generated JSON twins, malformed files).

    The repository treats YAML as the source of truth and generates the sibling ``.json``
    file with Rain. A JSON file whose YAML sibling exists is a *generated twin*: it is not
    assessed separately (that would double every finding) but it is compared with its
    source so that drift is reported. ``include_generated_json`` assesses twins anyway.

    A well-formed file that is not a CloudFormation template (a Kubernetes manifest, an event
    fixture) is *skipped*. A file with a template extension that is not well-formed JSON/YAML is
    *malformed*: nothing can tell whether it is a template, so it is reported as a coverage gap
    rather than silently dropped.
    """
    templates: list[Path] = []
    terraform: list[Path] = []
    skipped: list[str] = []
    malformed: list[str] = []
    twins: dict[Path, Path] = {}
    for root, dirs, files in os.walk(repo):
        dirs[:] = sorted(d for d in dirs if d not in SKIP_DIRS and not d.startswith("."))
        for fn in sorted(files):
            p = Path(root) / fn
            if p.suffix == ".tf":
                terraform.append(p)
                continue
            if p.suffix not in TEMPLATE_EXTS:
                continue
            if fn in {".cfnlintrc", "cfn-lint.yaml", "package.json", "package-lock.json"} or fn.startswith("."):
                continue
            rel = str(p.relative_to(repo).as_posix())
            try:
                docs = parse_documents(p)
            except TemplateParseError:
                malformed.append(rel)
                continue
            if len(docs) != 1 or not isinstance(docs[0], dict) or not is_cfn_template(docs[0]):
                skipped.append(rel)
                continue
            if p.suffix == ".json":
                src = next((p.with_suffix(ext) for ext in (".yaml", ".yml") if p.with_suffix(ext).is_file()), None)
                if src is not None:
                    twins[p] = src
                    if not include_generated_json:
                        continue
            templates.append(p)
    return templates, terraform, skipped, twins, malformed


def _canon(node: object) -> object:
    """Text rendering of an expression for *identity* purposes (bucket-name matching, IAM statement digests): scalars
    become strings so ``Ref: X`` in YAML and JSON digest identically regardless of scalar typing."""
    if isinstance(node, dict):
        return {str(k): _canon(v) for k, v in node.items()}
    if isinstance(node, list):
        return [_canon(v) for v in node]
    if isinstance(node, bool):
        return str(node).lower()
    if node is None:
        return ""
    return str(node)


def _twin_canon(node: object) -> object:
    """Plain-Python rendering of a template for *drift comparison*: line-aware containers become dict/list and every
    scalar keeps its type, so ``80`` vs ``"80"``, ``true`` vs ``"true"`` and a dropped value (``null``) vs ``""`` are
    reported as drift. The one equivalence applied is ``Fn::GetAZs`` with an empty string versus null: YAML ``!GetAZs ""``
    and Rain's JSON ``{"Fn::GetAZs": null}`` both denote the current Region."""
    if isinstance(node, dict):
        if set(node) == {"Fn::GetAZs"} and node["Fn::GetAZs"] in (None, ""):
            return {"Fn::GetAZs": ""}
        return {str(k): _twin_canon(v) for k, v in node.items()}
    if isinstance(node, list):
        return [_twin_canon(v) for v in node]
    return node


TWIN_RULE = Rule(
    id="CSA-CFG-002",
    title="Generated JSON template is out of sync with its YAML source",
    severity=CAT_III,
    nist=["CM-2", "CM-3", "CM-6"],
    area="validity",
    description="The repository generates JSON templates from YAML sources. This JSON file differs from its YAML sibling, so the two deployable artifacts do not implement the same configuration.",
    remediation="Regenerate the JSON rendering from the YAML source (scripts/create-json-single.sh) and add the regeneration step to the CI gate so that the artifacts cannot drift.",
    applies_to=["AWS::CloudFormation::Template"],
)


def check_generated_twins(repo: Path, twins: dict[Path, Path]) -> tuple[list[Finding], set[str], list[str]]:
    """Returns (drift findings, twins compared with their source, twins that could not be compared).

    A twin is compared only when both it and its YAML source parse; otherwise it is reported as not compared
    so that the coverage gate and the baseline carry-forward do not treat it as assessed."""
    out: list[Finding] = []
    compared: set[str] = set()
    failed: list[str] = []
    for gen, src in sorted(twins.items()):
        rel = str(gen.relative_to(repo).as_posix())
        g, s = load_template(gen), load_template(src)
        if g is None or s is None:
            failed.append(rel)
            continue
        compared.add(rel)
        if _twin_canon(g) == _twin_canon(s):
            continue
        text = gen.read_text(encoding="utf-8", errors="replace")
        ctx = TemplateCtx(path=gen, rel=rel, text=text, lines=text.splitlines(), data=g, service_dir=service_dir_of(rel))
        hit = Hit("(template)", "AWS::CloudFormation::Template", 1, f"content differs from {src.relative_to(repo).as_posix()}")
        out.append(finding_from_hit(TWIN_RULE, ctx, hit))
    return out, compared, failed


def assess_templates(repo: Path, templates: list[Path]) -> tuple[list[Finding], dict[tuple[str, str], str], Counter, Counter, set[str], dict[str, list[str]]]:
    """Returns (findings, resource types, type counts, hits per rule, templates the rule pack parsed,
    rules that raised per template). A rule that raises is recorded rather than re-raised so one rule bug does not
    abort the run, but the (template, rule) pair is a coverage gap: it fails ``--fail-on-incomplete`` and blocks the
    baseline carry-forward from marking that rule's findings on that template as remediated."""
    findings: list[Finding] = []
    resource_types: dict[tuple[str, str], str] = {}
    type_counts: Counter = Counter()
    rule_hits: Counter = Counter()
    parsed: set[str] = set()
    rule_failures: dict[str, list[str]] = {}
    seen_ids: set[str] = set()
    for t in templates:
        rel = str(t.relative_to(repo).as_posix())
        text = t.read_text(encoding="utf-8", errors="replace")
        data = load_template(t)
        if data is None:
            continue
        parsed.add(rel)
        ctx = TemplateCtx(path=t, rel=rel, text=text, lines=text.splitlines(), data=data, service_dir=service_dir_of(rel))
        for lid, res in ctx.resources.items():
            rt = rtype_of(res)
            resource_types[(rel, lid)] = rt
            type_counts[rt] += 1
        for r in RULES:
            assert r.check is not None
            try:
                hits = list(r.check(ctx))
            except Exception as e:  # noqa: BLE001 - a rule bug must not abort the assessment
                print(f"  [warn] rule {r.id} failed on {rel}: {e}", file=sys.stderr)
                rule_failures.setdefault(rel, []).append(r.id)
                continue
            for hit in hits:
                f = finding_from_hit(r, ctx, hit)
                if f.finding_id in seen_ids:
                    print(f"  [warn] rule {r.id} reported {rel}:{hit.logical_id} twice with the same key; the duplicate is dropped", file=sys.stderr)
                    continue
                seen_ids.add(f.finding_id)
                findings.append(f)
                rule_hits[r.id] += 1
    return findings, resource_types, type_counts, rule_hits, parsed, rule_failures


# --------------------------------------------------------------------------------------
# Summaries
# --------------------------------------------------------------------------------------


def family_of(control: str) -> str:
    return NIST_CONTROLS.get(control, ("Other", ""))[0]


def summarize(findings: list[Finding], templates: list[Path]) -> dict:
    open_ = [f for f in findings if f.is_open]
    by_sev = Counter(f.severity for f in open_)
    by_sev_all = Counter(f.severity for f in findings)
    by_dir = Counter(f.service_directory for f in open_)
    by_family: Counter = Counter()
    for f in open_:
        for fam in sorted({family_of(c) for c in f.nist_controls}):
            by_family[fam] += 1
    by_disp = Counter(f.disposition for f in findings)
    by_source = Counter(f.source for f in findings)
    return {
        "total_findings": len(findings),
        "open_findings": len(open_),
        "open_by_severity": {s: by_sev.get(s, 0) for s in (CAT_I, CAT_II, CAT_III)},
        "all_by_severity": {s: by_sev_all.get(s, 0) for s in (CAT_I, CAT_II, CAT_III)},
        "by_disposition": dict(by_disp),
        "by_source": dict(by_source),
        "open_by_service_directory": dict(sorted(by_dir.items(), key=lambda kv: (-kv[1], kv[0]))),
        "open_by_control_family": dict(sorted(by_family.items(), key=lambda kv: (-kv[1], kv[0]))),
        "templates_assessed": len(templates),
        "templates_with_open_findings": len({f.template_path for f in open_}),
    }


def control_coverage(findings: list[Finding], templates_assessed: int) -> list[dict]:
    rows = []
    for control, (family, title) in NIST_CONTROLS.items():
        fs = [f for f in findings if control in f.nist_controls]
        if not fs:
            continue
        rows.append({
            "control": control, "family": family, "title": title,
            "findings": len(fs), "open": sum(f.is_open for f in fs), "remediated": sum(f.disposition == "Remediated in PR" for f in fs),
            "cat_i_open": sum(f.is_open and f.severity == CAT_I for f in fs),
            "templates_with_findings": len({f.template_path for f in fs}), "templates_assessed": templates_assessed,
        })
    rows.sort(key=lambda r: (-r["open"], r["control"]))
    return rows


def git_info(repo: Path) -> tuple[str, str]:
    try:
        sha = run(["git", "rev-parse", "HEAD"], repo).stdout.strip()
        branch = run(["git", "rev-parse", "--abbrev-ref", "HEAD"], repo).stdout.strip()
        return sha or "unknown", branch or "unknown"
    except OSError:
        return "unknown", "unknown"


def template_revision(repo: Path, paths: list[str]) -> tuple[str, int]:
    """(last commit that changed any assessed template, number of assessed templates with uncommitted edits).

    HEAD identifies the checkout; this identifies the template set the findings describe, which does not move
    when only assessment artifacts are committed on top."""
    if not paths:
        return "unknown", 0
    try:
        last = run(["git", "log", "-1", "--format=%H", "--", *paths], repo).stdout.strip()
        dirty = run(["git", "status", "--porcelain", "--", *paths], repo).stdout.splitlines()
        return last or "unknown", len(dirty)
    except OSError:
        return "unknown", 0


# --------------------------------------------------------------------------------------
# XLSX
# --------------------------------------------------------------------------------------

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF")
SEV_FILL = {CAT_I: "F8CBAD", CAT_II: "FFE699", CAT_III: "E2EFDA"}
FINDINGS_COLUMNS = [
    ("Finding ID", 22, lambda f: f.finding_id), ("Template path", 40, lambda f: f.template_path),
    ("Resource logical ID", 26, lambda f: f.resource_logical_id), ("Resource type", 30, lambda f: f.resource_type),
    ("Finding title", 48, lambda f: f.title), ("Description", 70, lambda f: f.description),
    ("Evidence (file:line snippet)", 60, lambda f: f"{f.evidence_file}:{f.evidence_line}  {f.evidence_snippet}".strip()),
    ("Severity (CAT I/II/III)", 12, lambda f: f.severity), ("Severity justification", 40, lambda f: f.severity_justification),
    ("Risk rank (1-N)", 10, lambda f: f.risk_rank), ("Risk score", 10, lambda f: f.risk_score),
    ("NIST 800-53 Rev 5 control(s)", 24, lambda f: ", ".join(f.nist_controls)), ("CIS AWS v3.0 ID", 12, lambda f: f.cis_aws_v3_id or ""),
    ("DoD Cloud SRG area", 50, lambda f: f.dod_cloud_srg_area), ("Tool/rule ID", 28, lambda f: f.tool_rule_id),
    ("Recommended remediation", 70, lambda f: f.recommended_remediation), ("Disposition", 24, lambda f: f.disposition),
    ("Disposition note", 40, lambda f: f.disposition_note), ("Owner (role)", 30, lambda f: f.owner), ("Target date", 40, lambda f: f.target_date),
    ("Service directory", 16, lambda f: f.service_directory), ("Source", 14, lambda f: f.source), ("Evidence detail", 60, lambda f: f.evidence),
]


def style_header(ws, ncols: int) -> None:
    for c in range(1, ncols + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"


def add_table(ws, headers: list[str], rows: list[list[Any]], widths: list[int] | None = None, autofilter: bool = True) -> None:
    ws.append(headers)
    for r in rows:
        ws.append(r)
    style_header(ws, len(headers))
    for i, h in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(i)].width = (widths[i - 1] if widths else max(12, min(60, len(h) + 4)))
    if autofilter and rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"


def write_xlsx(path: Path, findings: list[Finding], summary: dict, before: dict | None, coverage: list[dict], meta: dict, tools: list[ToolResult]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Findings"
    headers = [c[0] for c in FINDINGS_COLUMNS]
    rows = [[c[2](f) for c in FINDINGS_COLUMNS] for f in findings]
    add_table(ws, headers, rows, [c[1] for c in FINDINGS_COLUMNS])
    for row in ws.iter_rows(min_row=2, max_row=len(rows) + 1):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    sev_col = get_column_letter(headers.index("Severity (CAT I/II/III)") + 1)
    last = len(rows) + 1
    if rows:
        rng = f"A2:{get_column_letter(len(headers))}{last}"
        for sev, color in SEV_FILL.items():
            ws.conditional_formatting.add(rng, FormulaRule(formula=[f'${sev_col}2="{sev}"'], fill=PatternFill("solid", fgColor=color), stopIfTrue=False))
        disp_col = get_column_letter(headers.index("Disposition") + 1)
        ws.conditional_formatting.add(rng, FormulaRule(formula=[f'${disp_col}2="Remediated in PR"'], font=Font(color="808080", strike=False),
                                                       fill=PatternFill("solid", fgColor="D9D9D9"), stopIfTrue=True))

    # Summary
    ws2 = wb.create_sheet("Summary")
    ws2.append(["Cloud Security Assessment — Summary"])
    ws2["A1"].font = Font(bold=True, size=14)
    ws2.append(["Scan date (UTC)", meta["scan_date"]])
    ws2.append(["Commit SHA", meta["commit_sha"]])
    ws2.append(["Branch", meta["branch"]])
    ws2.append(["Templates last changed in commit", meta["templates_commit_sha"]])
    ws2.append(["Templates with uncommitted edits at scan time", meta["templates_with_uncommitted_edits"]])
    ws2.append(["Templates assessed", summary["templates_assessed"]])
    ws2.append(["Terraform files assessed", meta["terraform_files"]])
    ws2.append(["Total findings (all dispositions)", summary["total_findings"]])
    ws2.append(["Open findings", summary["open_findings"]])
    ws2.append([])
    ws2.append(["Open findings by severity", "Count"] + (["Before remediation (baseline)", "After remediation (this run)", "Change"] if before else []))
    for sev in (CAT_I, CAT_II, CAT_III):
        row = [sev, summary["open_by_severity"][sev]]
        if before:
            b = before.get("open_by_severity", {}).get(sev, 0)
            row += [b, summary["open_by_severity"][sev], summary["open_by_severity"][sev] - b]
        ws2.append(row)
    total_row = ["Total", summary["open_findings"]]
    if before:
        b = before.get("open_findings", 0)
        total_row += [b, summary["open_findings"], summary["open_findings"] - b]
    ws2.append(total_row)
    ws2.append([])
    ws2.append(["Findings by disposition", "Count"])
    for d in DISPOSITIONS:
        ws2.append([d, summary["by_disposition"].get(d, 0)])
    ws2.append([])
    ws2.append(["Open findings by service directory", "Count"])
    for k, v in summary["open_by_service_directory"].items():
        ws2.append([k, v])
    ws2.append([])
    ws2.append(["Open findings by NIST control family", "Count"])
    for k, v in summary["open_by_control_family"].items():
        ws2.append([k, v])
    ws2.append([])
    ws2.append(["Findings by source", "Count"])
    for k, v in sorted(summary["by_source"].items()):
        ws2.append([k, v])
    for r in ws2.iter_rows():
        for c in r:
            if c.column == 1 and c.value in {"Open findings by severity", "Findings by disposition", "Open findings by service directory",
                                             "Open findings by NIST control family", "Findings by source"}:
                c.font = Font(bold=True)
    ws2.column_dimensions["A"].width = 44
    for col in "BCDE":
        ws2.column_dimensions[col].width = 30

    # Control coverage
    ws3 = wb.create_sheet("Control-Coverage")
    add_table(ws3, ["NIST 800-53 Rev 5 control", "Family", "Control title", "Findings (all)", "Open", "Remediated in PR", "Open CAT I",
                    "Templates with findings", "Templates assessed"],
              [[r["control"], r["family"], r["title"], r["findings"], r["open"], r["remediated"], r["cat_i_open"], r["templates_with_findings"], r["templates_assessed"]] for r in coverage],
              [26, 40, 60, 14, 10, 16, 12, 22, 18])

    # POA&M draft
    ws4 = wb.create_sheet("POAM-Draft")
    poam_rows = []
    for f in findings:
        if not f.is_open or f.severity not in (CAT_I, CAT_II):
            continue
        poam_rows.append([
            f"POAM-{f.risk_rank:04d}", f.finding_id, f"{f.title} — {f.template_path} [{f.resource_logical_id}]",
            ", ".join(f.nist_controls), f.severity, f.risk_rank, f"{f.evidence_file}:{f.evidence_line}",
            f.target_date,
            "M1: system owner confirms disposition; M2: template change reviewed and merged; M3: re-scan shows finding closed; M4: deployed stacks updated",
            f"{f.owner}; template author; ISSO review",
            "Ongoing", f.recommended_remediation, f.dod_cloud_srg_area, f.cis_aws_v3_id or "",
        ])
    add_table(ws4, ["POA&M ID", "Finding ID", "Weakness", "Control (NIST 800-53 Rev 5)", "Severity", "Risk rank", "Evidence", "Scheduled completion (placeholder)",
                    "Milestones", "Resources required", "Status", "Remediation", "DoD Cloud SRG area", "CIS AWS v3.0 ID"],
              poam_rows, [12, 22, 70, 24, 10, 10, 40, 40, 70, 40, 12, 70, 50, 12])
    if poam_rows:
        rng = f"A2:{get_column_letter(14)}{len(poam_rows) + 1}"
        for sev, color in SEV_FILL.items():
            ws4.conditional_formatting.add(rng, FormulaRule(formula=[f'$E2="{sev}"'], fill=PatternFill("solid", fgColor=color)))

    # Method
    ws5 = wb.create_sheet("Method")
    ws5.append(["Item", "Detail"])
    method_rows = [
        ["Scan date (UTC)", meta["scan_date"]], ["Commit SHA", meta["commit_sha"]], ["Branch", meta["branch"]],
        ["Templates last changed in commit", meta["templates_commit_sha"]],
        ["Templates with uncommitted edits at scan time", meta["templates_with_uncommitted_edits"]],
        ["Runner", f"scripts/cloud_security_assessment.py (Python {platform.python_version()}, PyYAML {yaml.__version__})"],
        ["Templates discovered", f"{summary['templates_assessed']} CloudFormation templates (.yaml/.yml/.json/.template with a Resources section)"],
        ["Files skipped (not CloudFormation)", str(meta["skipped_files"])],
        ["Malformed JSON/YAML files (not assessed by any scanner)", ", ".join(meta["malformed_files"]) or "none"],
        ["Custom rule failures (template: rules)", "; ".join(f"{p}: {', '.join(ids)}" for p, ids in sorted(meta["custom_rule_failures"].items())) or "none"],
        ["Terraform", f"{meta['terraform_files']} .tf files found. " + ("Scanned with Checkov (terraform framework)." if meta["terraform_files"] else
                      "No Terraform sources exist in this revision; the Terraform check is not applicable. The runner scans .tf files automatically when present.")],
    ]
    for t in tools:
        method_rows.append([f"Tool: {t.name}", f"version: {t.version}; status: {t.status}; raw results: {t.raw_count}; duration: {t.duration_s}s. {t.note}".strip()])
    method_rows += [
        ["Custom rule pack", f"{len(RULES) + 1} rules (CSA-*), including the generated-JSON drift check. Categories: encryption at rest, KMS key ownership, TLS in transit, public exposure, logging and monitoring, IAM least privilege, IMDSv2, secrets and credentials, backup/retention and deletion protection, template configuration drift."],
        ["Normalization", "Every tool result is mapped to one schema. A Checkov/cfn_nag result that reports the same weakness on the same resource as a custom rule is merged into that finding (tool IDs kept in Tool/rule ID); when the custom rule reported that resource several times (one finding per ingress rule or policy statement) the tool result attaches only to the finding for the same port or source line, otherwise it stays a tool finding. A cfn_nag result that repeats an equivalent Checkov check on the same resource is attached to the Checkov finding. Other tool results become their own findings with keyword-based control mapping."],
        ["Severity model", "DISA CAT I: direct and immediate loss of confidentiality, integrity or availability. CAT II: potential loss. CAT III: degraded protection. Each rule carries a default CAT that specific evidence may lower (for example an open CIDR that comes from a parameter default is CAT II rather than CAT I)."],
        ["Risk rank", "Rank 1 is highest risk. Score = severity base (CAT I 300 / CAT II 200 / CAT III 100) + exposure 40 + credential 30 + data store 20 + 10 per confirming external tool + 5 for custom-rule evidence + 3 for core service directories. Open findings rank before closed ones."],
        ["Control mappings", "NIST SP 800-53 Rev. 5 control IDs are assigned per rule. CIS AWS Foundations Benchmark v3.0.0 recommendation IDs are included only where the AWS Security Hub CIS v3.0.0 mapping confirms a matching control; otherwise the CIS field is blank. DoD Cloud Computing SRG areas use the SRG section names and Cloud Computing Mission Owner SRG requirement IDs."],
        ["cfn-lint handling", "Only Error-level results are findings (CM-2/CM-6, CAT III). Warning and informational messages are counted in the tool note. The repository .cfnlintrc (ignored checks and templates) is honored."],
        ["Severity policy", "CAT I is assigned only by the custom rule pack, which verifies the exact condition on the parsed template (for example an ingress rule to TCP/22 from 0.0.0.0/0 with no parameter override). Checkov and cfn_nag hits that confirm such a condition are merged into the custom finding and listed as corroborating tool IDs. Tool-only hits that rely on pattern heuristics (for example secret-like strings in user data, or wildcard detection that also matches scoped service wildcards) are capped at CAT II and carry the tool's own justification."],
        ["Known limits", "Static analysis of templates only: no deployed-account evidence, no AWS Config or Security Hub data. Parameter values are evaluated from Defaults; values supplied at deploy time are unknown. Intrinsic functions other than Ref/Fn::If defaults are treated as unresolved and never generate a finding on their own. AWS::LanguageExtensions Fn::ForEach loops over literal lists are expanded into their resources (a loop over a parameter is assessed once with the ${Identifier} placeholder in its logical ID); nested-stack and macro-generated resources are not expanded. Template intent (for example an Internet-facing web tier) is inferred, not confirmed."],
        ["Dispositions", "Open / Remediated in PR / Risk acceptance recommended / Not applicable. 'Remediated in PR' is assigned automatically when a baseline findings.json (--baseline) contains a finding that the current run no longer detects and every tool that produced it ran again on that template; when a tool was skipped, failed or could not process the template the baseline disposition is carried forward with a 'Not re-evaluated' note. Government overrides can be supplied with --dispositions."],
    ]
    for r in method_rows:
        ws5.append(r)
    style_header(ws5, 2)
    ws5.column_dimensions["A"].width = 34
    ws5.column_dimensions["B"].width = 140
    for row in ws5.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(path)


# --------------------------------------------------------------------------------------
# Main
# --------------------------------------------------------------------------------------


def console_summary(summary: dict, before: dict | None, tools: list[ToolResult], findings: list[Finding], out_dir: Path) -> None:
    print("\n=== Cloud Security Assessment summary ===")
    for t in tools:
        print(f"  {t.name:9s} {t.status:8s} {t.version[:40]:40s} raw={t.raw_count} {t.note[:80]}")
    print(f"  templates assessed: {summary['templates_assessed']}   findings: {summary['total_findings']} total / {summary['open_findings']} open")
    for sev in (CAT_I, CAT_II, CAT_III):
        line = f"  {sev:7s} open: {summary['open_by_severity'][sev]:4d}"
        if before:
            line += f"   (before: {before.get('open_by_severity', {}).get(sev, 0)})"
        print(line)
    print("  dispositions: " + ", ".join(f"{k}={v}" for k, v in sorted(summary["by_disposition"].items())))
    print("  top 10 open findings by risk rank:")
    for f in [f for f in findings if f.is_open][:10]:
        print(f"    #{f.risk_rank:<3d} {f.severity:7s} {f.finding_id:26s} {f.template_path}:{f.evidence_line} [{f.resource_logical_id}] {f.title[:60]}")
    print(f"  outputs: {out_dir / 'findings.json'}, {out_dir / 'Cloud-Security-Findings-Tracker.xlsx'}")


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--repo-root", default=".", help="repository root (default: current directory)")
    ap.add_argument("--out-dir", default="security-assessment", help="output directory relative to repo root")
    ap.add_argument("--skip-checkov", action="store_true")
    ap.add_argument("--skip-cfn-nag", action="store_true")
    ap.add_argument("--skip-cfn-lint", action="store_true")
    ap.add_argument("--baseline", help="previous findings.json (path relative to the current directory, not --repo-root); "
                                        "findings that disappear are recorded as 'Remediated in PR'")
    ap.add_argument("--dispositions", help="JSON file {finding_id: {disposition, note}} with Government disposition overrides "
                                            "(path relative to the current directory, not --repo-root)")
    ap.add_argument("--include-generated-json", action="store_true",
                    help="also assess .json templates that have a YAML sibling (by default they are only checked for drift)")
    ap.add_argument("--fail-on-incomplete", action="store_true",
                    help="exit 2 when a scanner did not run, failed, or left templates unprocessed (for CI gates)")
    args = ap.parse_args(argv)

    repo = Path(args.repo_root).resolve()
    out_dir = (repo / args.out_dir).resolve()
    out_dir.mkdir(parents=True, exist_ok=True)
    sha, branch = git_info(repo)
    scan_date = dt.datetime.now(dt.timezone.utc).strftime("%Y-%m-%d %H:%M UTC")

    print(f"Repository: {repo}\nCommit: {sha} ({branch})")
    templates, terraform, skipped, twins, malformed = discover_templates(repo, args.include_generated_json)
    print(f"Discovered {len(templates)} CloudFormation templates, {len(terraform)} Terraform files; "
          f"{len(twins)} generated JSON twins of YAML sources; skipped {len(skipped)} non-template files; "
          f"{len(malformed)} malformed JSON/YAML file(s)")

    print("Running custom rule pack ...")
    findings, resource_types, type_counts, rule_hits, parsed, rule_failures = assess_templates(repo, templates)
    twin_findings, twins_compared, twins_failed = check_generated_twins(repo, twins)
    findings += twin_findings
    parsed |= twins_compared
    rule_hits[TWIN_RULE.id] += len(twin_findings)
    print(f"  custom rule findings: {len(findings)} ({len(twin_findings)} generated JSON files differ from their YAML source; "
          f"{len(twins_failed)} twin(s) could not be compared)")

    tools: list[ToolResult] = []
    external: list[ExternalHit] = []
    tr = ToolResult("checkov", "n/a", "skipped", "skipped by --skip-checkov")
    if not args.skip_checkov:
        print("Running checkov ...")
        external += run_checkov(repo, templates, terraform, tr)
        print(f"  checkov: {tr.status} ({tr.raw_count} failed checks) {tr.note}")
    tools.append(tr)
    tr = ToolResult("cfn_nag", "n/a", "skipped", "skipped by --skip-cfn-nag")
    if not args.skip_cfn_nag:
        print("Running cfn_nag ...")
        external += run_cfn_nag(repo, templates, tr)
        print(f"  cfn_nag: {tr.status} ({tr.raw_count} violations) {tr.note}")
    tools.append(tr)
    tr = ToolResult("cfn-lint", "n/a", "skipped", "skipped by --skip-cfn-lint")
    lint_warnings = 0
    if not args.skip_cfn_lint:
        print("Running cfn-lint ...")
        lint_hits, lint_warnings = run_cfn_lint(repo, templates, tr)
        external += lint_hits
        print(f"  cfn-lint: {tr.status} ({tr.raw_count} errors, {lint_warnings} warnings/informational)")
    tools.append(tr)

    assessed = {str(t.relative_to(repo).as_posix()) for t in templates} | {str(t.relative_to(repo).as_posix()) for t in terraform}
    templates_sha, templates_dirty = template_revision(repo, sorted(assessed))
    findings = merge_external(findings, external, repo, resource_types, assessed)
    findings, before = apply_baseline(findings, Path(args.baseline) if args.baseline else None, Path(args.dispositions) if args.dispositions else None,
                                      tools, parsed, rule_failures)
    rank(findings)
    summary = summarize(findings, templates)
    coverage = control_coverage(findings, len(templates))

    meta = {
        "scan_date": scan_date, "commit_sha": sha, "branch": branch, "repo_root": repo.name,
        "templates_commit_sha": templates_sha, "templates_with_uncommitted_edits": templates_dirty,
        "templates_assessed": len(templates), "terraform_files": len(terraform), "skipped_files": len(skipped),
        "malformed_files": malformed, "custom_rule_failures": rule_failures,
        "generated_json_twins": {str(g.relative_to(repo).as_posix()): str(y.relative_to(repo).as_posix()) for g, y in sorted(twins.items())},
        "generated_json_assessed": args.include_generated_json,
        "generated_json_not_compared": twins_failed,
        "tools": [t.__dict__ for t in tools],
        "custom_rules": [{"id": r.id, "title": r.title, "default_severity": r.severity, "nist": r.nist, "cis_aws_v3_id": r.cis,
                          "dod_cloud_srg_area": SRG_AREAS[r.area], "applies_to": r.applies_to, "overlapping_tool_ids": sorted(r.overlaps),
                          "hits": rule_hits.get(r.id, 0)} for r in [*RULES, TWIN_RULE]],
        "resource_type_counts": dict(type_counts.most_common()),
        "template_paths": [str(t.relative_to(repo).as_posix()) for t in templates],
        "terraform_paths": [str(t.relative_to(repo).as_posix()) for t in terraform],
        "baseline": Path(args.baseline).name if args.baseline else None, "cfn_lint_warnings": lint_warnings,
    }
    payload = {"metadata": meta, "summary": summary, "before_remediation": before, "control_coverage": coverage,
               "findings": [f.to_dict() for f in findings]}
    (out_dir / "findings.json").write_text(json.dumps(payload, indent=2, default=str) + "\n")
    write_xlsx(out_dir / "Cloud-Security-Findings-Tracker.xlsx", findings, summary, before, coverage, meta, tools)
    console_summary(summary, before, tools, findings, out_dir)
    incomplete = [f"{t.name}: {t.status}" for t in tools if t.status != "ran"] + [f"{t.name}: {len(t.unparsed)} template(s) not processed" for t in tools if t.unparsed]
    unparsed_custom = sorted(t for t in meta["template_paths"] if t not in parsed)
    if unparsed_custom:
        incomplete.append(f"custom rule pack: {len(unparsed_custom)} template(s) could not be parsed: " + ", ".join(unparsed_custom))
    if rule_failures:
        incomplete.append("custom rule pack: rule failures on " + "; ".join(f"{p} ({', '.join(ids)})" for p, ids in sorted(rule_failures.items())))
    if malformed:
        incomplete.append(f"discovery: {len(malformed)} file(s) with a template extension are not well-formed JSON/YAML and were not assessed by any scanner: " + ", ".join(malformed))
    if twins_failed:
        incomplete.append(f"generated JSON: {len(twins_failed)} twin(s) could not be compared with their YAML source: " + ", ".join(twins_failed))
    if incomplete:
        print("Coverage gaps: " + "; ".join(incomplete))
        if args.fail_on_incomplete:
            return 2
    return 0


if __name__ == "__main__":
    sys.exit(main())
